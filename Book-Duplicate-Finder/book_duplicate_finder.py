"""Book / Document Content Duplicate Finder.

Beginner summary
----------------
This script walks the folders listed in ``search_paths.txt`` and looks for
books/documents that appear to contain the same material even when the files
have different names, sizes, page counts, compression, or formats.

It does NOT delete, move, rename, or modify your books. It only creates reports.

How matching works, in plain English:
1. Exact SHA-256 hashes catch byte-for-byte duplicates.
2. Text is extracted and normalized so formatting differences matter less.
3. Small overlapping word sequences ("shingles") are hashed into a compact
   content sketch. Similar sketches suggest similar content.
4. A containment score estimates how much of the smaller document appears in
   the larger one, useful when one copy has extra pages.
5. For image-heavy PDFs/DJVU/CBZ/CBR, sampled pages are compared with
   perceptual image hashes.
6. Likely matches are grouped and written to HTML/CSV for human review.

Configuration lives in ``settings.ini``. Runtime caches and reports are stored
under ``cache`` and ``results`` respectively.
"""

from __future__ import annotations

import configparser
import csv
import hashlib
import html
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import webbrowser
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

ROOT = Path(__file__).resolve().parent
DEPS = ROOT / "deps"
if DEPS.exists():
    sys.path.insert(0, str(DEPS))

try:
    import pymupdf as fitz
except Exception:
    import fitz  # type: ignore

from bs4 import BeautifulSoup
from ebooklib import ITEM_DOCUMENT, epub
from PIL import Image
import imagehash
from docx import Document
from rapidfuzz.fuzz import ratio as fuzz_ratio
from striprtf.striprtf import rtf_to_text

try:
    import mobi
except Exception:
    mobi = None

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
CONFIG = configparser.ConfigParser()
CONFIG.read(ROOT / "settings.ini", encoding="utf-8")

EXTENSIONS = {x.strip().lower() for x in CONFIG.get("scan", "extensions").split(",") if x.strip()}
MIN_FILE_SIZE = CONFIG.getint("scan", "min_file_size", fallback=1024)
MAX_TEXT_CHARS = CONFIG.getint("scan", "max_text_chars", fallback=4_000_000)
SKETCH_SIZE = CONFIG.getint("fingerprint", "sketch_size", fallback=512)
SHINGLE_WORDS = CONFIG.getint("fingerprint", "shingle_words", fallback=5)
SHINGLE_STEP = CONFIG.getint("fingerprint", "shingle_step", fallback=2)
IMAGE_SAMPLES = CONFIG.getint("fingerprint", "image_samples", fallback=12)
LOW_TEXT_CHARS = CONFIG.getint("fingerprint", "low_text_chars", fallback=2500)

HIGH_CONTAIN = CONFIG.getfloat("matching", "high_containment", fallback=0.965)
HIGH_JACC = CONFIG.getfloat("matching", "high_jaccard", fallback=0.72)
PROB_CONTAIN = CONFIG.getfloat("matching", "probable_containment", fallback=0.90)
PROB_JACC = CONFIG.getfloat("matching", "probable_jaccard", fallback=0.58)
REVIEW_CONTAIN = CONFIG.getfloat("matching", "review_containment", fallback=0.80)
REVIEW_JACC = CONFIG.getfloat("matching", "review_jaccard", fallback=0.45)
HIGH_IMG = CONFIG.getfloat("matching", "high_image_similarity", fallback=0.90)
PROB_IMG = CONFIG.getfloat("matching", "probable_image_similarity", fallback=0.84)
REVIEW_IMG = CONFIG.getfloat("matching", "review_image_similarity", fallback=0.78)
MIN_SHARED = CONFIG.getint("matching", "minimum_shared_sketch_hashes", fallback=2)
OPEN_REPORT = CONFIG.getboolean("output", "open_report_after_scan", fallback=True)

CACHE_DIR = ROOT / "cache"
RESULTS_DIR = ROOT / "results"
CACHE_DB = CACHE_DIR / "fingerprints.sqlite"
CACHE_DIR.mkdir(exist_ok=True)
RESULTS_DIR.mkdir(exist_ok=True)

TEXT_EXTS = {".txt", ".html", ".htm", ".xhtml", ".xml", ".fb2"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}


# ---------------------------------------------------------------------------
# Folder discovery
# ---------------------------------------------------------------------------
def read_path_file(name: str) -> list[Path]:
    out: list[Path] = []
    p = ROOT / name
    if not p.exists():
        return out
    for raw in p.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        s = raw.strip().strip('"')
        if not s or s.startswith("#"):
            continue
        out.append(Path(os.path.expandvars(os.path.expanduser(s))).resolve())
    return out


def is_within(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except Exception:
        return False


def scan_files(search_roots: list[Path], excludes: list[Path]) -> list[Path]:
    found: list[Path] = []
    seen: set[str] = set()
    for root in search_roots:
        if not root.exists():
            print(f"WARNING: Search path does not exist: {root}")
            continue
        for base, dirs, files in os.walk(root):
            base_p = Path(base)
            dirs[:] = [d for d in dirs if not any(is_within(base_p / d, ex) for ex in excludes)]
            if any(is_within(base_p, ex) for ex in excludes):
                continue
            for fn in files:
                p = base_p / fn
                if p.suffix.lower() not in EXTENSIONS:
                    continue
                try:
                    if p.stat().st_size < MIN_FILE_SIZE:
                        continue
                    key = str(p.resolve()).lower()
                    if key not in seen:
                        seen.add(key)
                        found.append(p.resolve())
                except OSError:
                    pass
    found.sort(key=lambda x: str(x).lower())
    return found


def find_local_tool(filename: str) -> Path | None:
    for p in (ROOT / "tools").rglob(filename):
        if p.is_file():
            return p
    for p in (ROOT / "tools").rglob("*"):
        if p.is_file() and p.name.lower() == filename.lower():
            return p
    return None


def find_archive_tool() -> Path | None:
    # Full 7z.exe supports CHM/RAR-family extraction through its local DLLs.
    return find_local_tool("7z.exe")


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(4 * 1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_text(text: str) -> str:
    text = text.replace("\x00", " ").lower()
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"\b(?:page|pg)\s*\d+\b", " ", text)
    text = re.sub(r"\b\d{1,4}\s*/\s*\d{1,4}\b", " ", text)
    text = re.sub(r"[^\w\s']+", " ", text, flags=re.UNICODE)
    text = re.sub(r"\b\d+\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    if MAX_TEXT_CHARS > 0 and len(text) > MAX_TEXT_CHARS:
        # Preserve the beginning, middle and end rather than only truncating the tail.
        third = MAX_TEXT_CHARS // 3
        mid = len(text) // 2
        text = text[:third] + " " + text[mid - third // 2: mid + third // 2] + " " + text[-third:]
    return text


def clean_title(path: Path) -> str:
    s = path.stem.lower()
    s = re.sub(r"\b(copy|final|new|old|scan|scanned|ocr|ebook|book|v\d+|rev\d*|edition|ed)\b", " ", s)
    s = re.sub(r"\b\d{4}\b", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def token_set_title(path: Path) -> set[str]:
    return {t for t in clean_title(path).split() if len(t) >= 3}


def soup_text(blob: bytes | str) -> str:
    try:
        soup = BeautifulSoup(blob, "lxml")
    except Exception:
        soup = BeautifulSoup(blob, "html.parser")
    for tag in soup(["script", "style", "nav"]):
        tag.decompose()
    return soup.get_text(" ", strip=True)


def sample_indices(n: int, count: int) -> list[int]:
    if n <= 0:
        return []
    count = max(1, min(count, n))
    if count == 1:
        return [0]
    vals = {round(i * (n - 1) / (count - 1)) for i in range(count)}
    return sorted(vals)


def phash_pil(img: Image.Image) -> str:
    if img.mode not in ("L", "RGB"):
        img = img.convert("RGB")
    return str(imagehash.phash(img, hash_size=8))


# ---------------------------------------------------------------------------
# Format-specific content extraction
# ---------------------------------------------------------------------------
def extract_pdf(path: Path) -> tuple[str, int | None, list[str]]:
    doc = fitz.open(str(path))
    texts: list[str] = []
    for page in doc:
        try:
            texts.append(page.get_text("text"))
        except Exception:
            texts.append("")
    text = "\n".join(texts)
    hashes: list[str] = []
    if len(normalize_text(text)) < LOW_TEXT_CHARS:
        for i in sample_indices(len(doc), IMAGE_SAMPLES):
            try:
                page = doc.load_page(i)
                pix = page.get_pixmap(matrix=fitz.Matrix(0.8, 0.8), alpha=False)
                mode = "RGB" if pix.n >= 3 else "L"
                img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                hashes.append(phash_pil(img))
            except Exception:
                continue
    pages = len(doc)
    doc.close()
    return text, pages, hashes


def extract_epub(path: Path) -> tuple[str, int | None, list[str]]:
    book = epub.read_epub(str(path), options={"ignore_ncx": True})
    parts = []
    for item in book.get_items():
        if item.get_type() == ITEM_DOCUMENT:
            parts.append(soup_text(item.get_content()))
    return "\n".join(parts), None, []


def extract_docx(path: Path) -> tuple[str, int | None, list[str]]:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts), None, []


def extract_text_file(path: Path) -> tuple[str, int | None, list[str]]:
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    if path.suffix.lower() in {".html", ".htm", ".xhtml", ".xml", ".fb2"}:
        text = soup_text(text)
    return text, None, []


def extract_rtf(path: Path) -> tuple[str, int | None, list[str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    return rtf_to_text(raw), None, []


def extract_chm(path: Path) -> tuple[str, int | None, list[str]]:
    seven = find_archive_tool()
    if not seven:
        raise RuntimeError("Local full 7z.exe not found. Run RUN.bat again.")
    with tempfile.TemporaryDirectory(prefix="bookdup_chm_") as td:
        out = Path(td)
        cp = subprocess.run([str(seven), "x", str(path), f"-o{out}", "-y"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=180)
        if cp.returncode not in (0, 1):
            raise RuntimeError(f"7-Zip CHM extraction failed with code {cp.returncode}")
        parts: list[str] = []
        for ext in ("*.htm", "*.html", "*.xhtml", "*.txt"):
            for f in out.rglob(ext):
                try:
                    if f.stat().st_size > 10 * 1024 * 1024:
                        continue
                    blob = f.read_bytes()
                    parts.append(soup_text(blob) if f.suffix.lower() != ".txt" else blob.decode("utf-8", errors="replace"))
                except Exception:
                    continue
        return "\n".join(parts), None, []


def djvu_page_count(path: Path) -> int | None:
    tool = find_local_tool("djvused.exe")
    if not tool:
        return None
    try:
        cp = subprocess.run([str(tool), "-e", "n", str(path)], capture_output=True, text=True, timeout=60)
        m = re.search(r"\d+", cp.stdout)
        return int(m.group()) if m else None
    except Exception:
        return None


def extract_djvu(path: Path) -> tuple[str, int | None, list[str]]:
    txt_tool = find_local_tool("djvutxt.exe")
    if not txt_tool:
        raise RuntimeError("Local djvutxt.exe not found. Run RUN.bat again.")
    cp = subprocess.run([str(txt_tool), str(path)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=240)
    text = cp.stdout.decode("utf-8", errors="replace") if cp.stdout else ""
    pages = djvu_page_count(path)
    hashes: list[str] = []
    ddjvu = find_local_tool("ddjvu.exe")
    if ddjvu and pages and len(normalize_text(text)) < LOW_TEXT_CHARS:
        with tempfile.TemporaryDirectory(prefix="bookdup_djvu_") as td:
            for idx0 in sample_indices(pages, IMAGE_SAMPLES):
                page_num = idx0 + 1
                out = Path(td) / f"p{page_num}.ppm"
                try:
                    cp2 = subprocess.run([str(ddjvu), "-format=ppm", f"-page={page_num}", str(path), str(out)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
                    if cp2.returncode == 0 and out.exists():
                        with Image.open(out) as img:
                            hashes.append(phash_pil(img))
                except Exception:
                    continue
    return text, pages, hashes


def extract_archive_images(path: Path) -> tuple[str, int | None, list[str]]:
    with tempfile.TemporaryDirectory(prefix="bookdup_comic_") as td:
        out = Path(td)
        if path.suffix.lower() == ".cbz":
            import zipfile
            with zipfile.ZipFile(path, "r") as z:
                z.extractall(out)
        else:
            seven = find_archive_tool()
            if not seven:
                raise RuntimeError("Local full 7z.exe not found for CBR.")
            cp = subprocess.run([str(seven), "x", str(path), f"-o{out}", "-y"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=240)
            if cp.returncode not in (0, 1):
                raise RuntimeError(f"CBR extraction failed with code {cp.returncode}")
        imgs = sorted([p for p in out.rglob("*") if p.is_file() and p.suffix.lower() in IMAGE_EXTS], key=lambda x: str(x).lower())
        hashes = []
        for i in sample_indices(len(imgs), IMAGE_SAMPLES):
            try:
                with Image.open(imgs[i]) as img:
                    hashes.append(phash_pil(img))
            except Exception:
                pass
        return "", len(imgs), hashes


def extract_mobi(path: Path) -> tuple[str, int | None, list[str]]:
    if mobi is None:
        raise RuntimeError("Python mobi package unavailable.")
    tempdir = None
    try:
        tempdir, extracted = mobi.extract(str(path))
        base = Path(tempdir)
        parts = []
        candidates = list(base.rglob("*.html")) + list(base.rglob("*.htm")) + list(base.rglob("*.xhtml")) + list(base.rglob("*.txt"))
        for f in candidates:
            try:
                b = f.read_bytes()
                parts.append(soup_text(b) if f.suffix.lower() != ".txt" else b.decode("utf-8", errors="replace"))
            except Exception:
                continue
        return "\n".join(parts), None, []
    finally:
        if tempdir:
            shutil.rmtree(tempdir, ignore_errors=True)


def extract_content(path: Path) -> tuple[str, int | None, list[str], str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        t, p, i = extract_pdf(path); return t, p, i, "pdf"
    if ext == ".epub":
        t, p, i = extract_epub(path); return t, p, i, "epub"
    if ext == ".docx":
        t, p, i = extract_docx(path); return t, p, i, "docx"
    if ext in TEXT_EXTS:
        t, p, i = extract_text_file(path); return t, p, i, "text"
    if ext == ".rtf":
        t, p, i = extract_rtf(path); return t, p, i, "rtf"
    if ext == ".chm":
        t, p, i = extract_chm(path); return t, p, i, "chm"
    if ext in {".djvu", ".djv"}:
        t, p, i = extract_djvu(path); return t, p, i, "djvu"
    if ext in {".cbz", ".cbr"}:
        t, p, i = extract_archive_images(path); return t, p, i, "comic"
    if ext in {".mobi", ".azw", ".azw3"}:
        t, p, i = extract_mobi(path); return t, p, i, "mobi"
    raise RuntimeError(f"No extractor for {ext}")


# ---------------------------------------------------------------------------
# Content fingerprinting
# ---------------------------------------------------------------------------
def hash64(data: bytes) -> int:
    return int.from_bytes(hashlib.blake2b(data, digest_size=8).digest(), "big", signed=False)


def content_sketch(text: str) -> tuple[list[int], int, int]:
    words = re.findall(r"[\w']+", text, flags=re.UNICODE)
    word_count = len(words)
    if len(words) < SHINGLE_WORDS:
        return [], word_count, 0
    vals: set[int] = set()
    for i in range(0, len(words) - SHINGLE_WORDS + 1, SHINGLE_STEP):
        sh = " ".join(words[i:i + SHINGLE_WORDS]).encode("utf-8", errors="ignore")
        vals.add(hash64(sh))
    unique_count = len(vals)
    if len(vals) <= SKETCH_SIZE:
        sketch = sorted(vals)
    else:
        import heapq
        sketch = sorted(heapq.nsmallest(SKETCH_SIZE, vals))
    return sketch, word_count, unique_count


# ---------------------------------------------------------------------------
# Cached fingerprint model
# ---------------------------------------------------------------------------
@dataclass
class Fingerprint:
    path: str
    size: int
    mtime_ns: int
    ext: str
    sha256: str
    text_hash: str
    word_count: int
    unique_shingles: int
    sketch: list[int]
    pages: int | None
    image_hashes: list[str]
    title: str
    extractor: str
    error: str = ""


def db_connect() -> sqlite3.Connection:
    con = sqlite3.connect(CACHE_DB)
    con.execute("""
        CREATE TABLE IF NOT EXISTS fp (
            path TEXT PRIMARY KEY,
            size INTEGER NOT NULL,
            mtime_ns INTEGER NOT NULL,
            payload TEXT NOT NULL
        )
    """)
    return con


def fp_to_dict(fp: Fingerprint) -> dict:
    return fp.__dict__.copy()


def fp_from_dict(d: dict) -> Fingerprint:
    return Fingerprint(**d)


def get_cached(con: sqlite3.Connection, path: Path) -> Fingerprint | None:
    try:
        st = path.stat()
    except OSError:
        return None
    row = con.execute("SELECT size, mtime_ns, payload FROM fp WHERE path=?", (str(path),)).fetchone()
    if not row:
        return None
    if row[0] != st.st_size or row[1] != st.st_mtime_ns:
        return None
    try:
        return fp_from_dict(json.loads(row[2]))
    except Exception:
        return None


def save_cached(con: sqlite3.Connection, fp: Fingerprint) -> None:
    con.execute(
        "INSERT OR REPLACE INTO fp(path,size,mtime_ns,payload) VALUES(?,?,?,?)",
        (fp.path, fp.size, fp.mtime_ns, json.dumps(fp_to_dict(fp), ensure_ascii=False)),
    )


def build_fp(path: Path) -> Fingerprint:
    st = path.stat()
    sha = file_sha256(path)
    try:
        text, pages, image_hashes, extractor = extract_content(path)
        norm = normalize_text(text)
        sketch, word_count, unique_shingles = content_sketch(norm)
        text_hash = hashlib.sha256(norm.encode("utf-8", errors="ignore")).hexdigest() if norm else ""
        return Fingerprint(
            path=str(path), size=st.st_size, mtime_ns=st.st_mtime_ns, ext=path.suffix.lower(),
            sha256=sha, text_hash=text_hash, word_count=word_count,
            unique_shingles=unique_shingles, sketch=sketch, pages=pages,
            image_hashes=image_hashes, title=clean_title(path), extractor=extractor,
        )
    except Exception as exc:
        return Fingerprint(
            path=str(path), size=st.st_size, mtime_ns=st.st_mtime_ns, ext=path.suffix.lower(),
            sha256=sha, text_hash="", word_count=0, unique_shingles=0, sketch=[], pages=None,
            image_hashes=[], title=clean_title(path), extractor="error", error=f"{type(exc).__name__}: {exc}",
        )


# ---------------------------------------------------------------------------
# Similarity calculations and candidate generation
# ---------------------------------------------------------------------------
def sketch_jaccard(a: Fingerprint, b: Fingerprint) -> float:
    if not a.sketch or not b.sketch:
        return 0.0
    k = min(SKETCH_SIZE, len(a.sketch), len(b.sketch))
    union_sorted = sorted(set(a.sketch).union(b.sketch))[:k]
    if not union_sorted:
        return 0.0
    aset, bset = set(a.sketch), set(b.sketch)
    both = sum(1 for x in union_sorted if x in aset and x in bset)
    return max(0.0, min(1.0, both / len(union_sorted)))


def containment_from_jaccard(j: float, na: int, nb: int) -> float:
    if j <= 0 or na <= 0 or nb <= 0:
        return 0.0
    inter = j * (na + nb) / (1.0 + j)
    return max(0.0, min(1.0, inter / min(na, nb)))


def phash_distance_hex(a: str, b: str) -> int:
    return bin(int(a, 16) ^ int(b, 16)).count("1")


def image_similarity(a: Fingerprint, b: Fingerprint) -> float:
    if not a.image_hashes or not b.image_hashes:
        return 0.0
    # For each image in the smaller sample set, find the closest page image in the other.
    x, y = (a.image_hashes, b.image_hashes) if len(a.image_hashes) <= len(b.image_hashes) else (b.image_hashes, a.image_hashes)
    scores = []
    for h in x:
        d = min(phash_distance_hex(h, h2) for h2 in y)
        scores.append(1.0 - d / 64.0)
    if not scores:
        return 0.0
    # Discount accidental blank/common-page matches by using both average and median-ish center.
    scores.sort()
    avg = sum(scores) / len(scores)
    med = scores[len(scores) // 2]
    return max(0.0, min(1.0, 0.6 * avg + 0.4 * med))


def title_similarity(a: Fingerprint, b: Fingerprint) -> float:
    if not a.title or not b.title:
        return 0.0
    return fuzz_ratio(a.title, b.title) / 100.0


def build_candidates(fps: list[Fingerprint]) -> set[tuple[int, int]]:
    candidates: set[tuple[int, int]] = set()

    # Exact hashes.
    by_sha: dict[str, list[int]] = defaultdict(list)
    for i, f in enumerate(fps):
        by_sha[f.sha256].append(i)
    for idxs in by_sha.values():
        if len(idxs) > 1:
            for x in range(len(idxs)):
                for y in range(x + 1, len(idxs)):
                    candidates.add((idxs[x], idxs[y]))

    # Shared content sketch values.
    inv: dict[int, list[int]] = defaultdict(list)
    prefix_n = min(160, SKETCH_SIZE)
    for i, f in enumerate(fps):
        for h in f.sketch[:prefix_n]:
            inv[h].append(i)
    pair_counts: Counter[tuple[int, int]] = Counter()
    for idxs in inv.values():
        if len(idxs) > 80:
            continue  # Extremely common hash: not selective enough.
        for x in range(len(idxs)):
            for y in range(x + 1, len(idxs)):
                a, b = idxs[x], idxs[y]
                if a > b:
                    a, b = b, a
                pair_counts[(a, b)] += 1
    for pair, cnt in pair_counts.items():
        if cnt >= MIN_SHARED:
            candidates.add(pair)

    # Similar filenames/titles + roughly plausible length/page count.
    token_inv: dict[str, list[int]] = defaultdict(list)
    for i, f in enumerate(fps):
        for t in token_set_title(Path(f.path)):
            token_inv[t].append(i)
    title_pair_counts: Counter[tuple[int, int]] = Counter()
    for idxs in token_inv.values():
        if len(idxs) > 120:
            continue
        for x in range(len(idxs)):
            for y in range(x + 1, len(idxs)):
                a, b = sorted((idxs[x], idxs[y]))
                title_pair_counts[(a, b)] += 1
    for pair, cnt in title_pair_counts.items():
        a, b = (fps[pair[0]], fps[pair[1]])
        if cnt >= 2 or title_similarity(a, b) >= 0.80:
            if a.word_count and b.word_count:
                r = min(a.word_count, b.word_count) / max(a.word_count, b.word_count)
                if r >= 0.55:
                    candidates.add(pair)
            elif a.pages and b.pages:
                r = min(a.pages, b.pages) / max(a.pages, b.pages)
                if r >= 0.55:
                    candidates.add(pair)
            elif a.image_hashes and b.image_hashes:
                candidates.add(pair)

    return candidates


# A Match is a pair that crossed one of the configured review thresholds.
@dataclass
class Match:
    a: int
    b: int
    classification: str
    jaccard: float
    containment: float
    image_similarity: float
    title_similarity: float
    exact: bool


def classify_pair(a: Fingerprint, b: Fingerprint) -> Match | None:
    exact = a.sha256 == b.sha256
    if exact:
        return Match(-1, -1, "EXACT", 1.0, 1.0, 1.0 if a.image_hashes and b.image_hashes else 0.0, title_similarity(a, b), True)

    j = sketch_jaccard(a, b)
    cont = containment_from_jaccard(j, a.unique_shingles, b.unique_shingles)
    img = image_similarity(a, b)
    ts = title_similarity(a, b)

    # Text is primary; image match covers scans/comics.
    if (cont >= HIGH_CONTAIN and j >= HIGH_JACC) or img >= HIGH_IMG:
        cls = "HIGH"
    elif (cont >= PROB_CONTAIN and j >= PROB_JACC) or img >= PROB_IMG:
        cls = "PROBABLE"
    elif (cont >= REVIEW_CONTAIN and j >= REVIEW_JACC) or img >= REVIEW_IMG:
        cls = "REVIEW"
    else:
        # Rescue strongly matching titles only when content is still meaningfully similar.
        if ts >= 0.93 and ((cont >= 0.72 and j >= 0.34) or img >= 0.72):
            cls = "REVIEW"
        else:
            return None
    return Match(-1, -1, cls, j, cont, img, ts, False)


class DSU:
    def __init__(self, n: int):
        self.p = list(range(n))
    def find(self, x: int) -> int:
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x
    def union(self, a: int, b: int) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.p[rb] = ra


def fmt_size(n: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    v = float(n)
    for u in units:
        if v < 1024 or u == units[-1]:
            return f"{v:.1f} {u}" if u != "B" else f"{int(v)} B"
        v /= 1024
    return f"{n} B"


def extra_hint(a: Fingerprint, b: Fingerprint) -> str:
    parts = []
    if a.pages and b.pages and a.pages != b.pages:
        bigger = a if a.pages > b.pages else b
        smaller = b if bigger is a else a
        parts.append(f"{Path(bigger.path).name} has {bigger.pages - smaller.pages:+d} more pages")
    if a.word_count and b.word_count:
        bigger = a if a.word_count > b.word_count else b
        smaller = b if bigger is a else a
        diff = bigger.word_count - smaller.word_count
        pct = diff / max(1, smaller.word_count) * 100
        if pct >= 1.0:
            parts.append(f"{Path(bigger.path).name} has about {pct:.1f}% more extracted words")
    return "; ".join(parts) if parts else "No clear page/text-length difference"


def esc(s: object) -> str:
    return html.escape(str(s), quote=True)


# ---------------------------------------------------------------------------
# Human-readable reports
# ---------------------------------------------------------------------------
def write_reports(fps: list[Fingerprint], matches: list[Match], groups: list[list[int]], elapsed: float) -> tuple[Path, Path, Path]:
    stamp = time.strftime("%Y%m%d_%H%M%S")
    html_path = RESULTS_DIR / f"duplicate_report_{stamp}.html"
    csv_path = RESULTS_DIR / f"duplicate_pairs_{stamp}.csv"
    err_path = RESULTS_DIR / f"scan_errors_{stamp}.csv"

    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["classification", "containment", "jaccard", "image_similarity", "title_similarity", "file_a", "file_b", "size_a", "size_b", "pages_a", "pages_b", "extra_hint"])
        for m in matches:
            a, b = fps[m.a], fps[m.b]
            w.writerow([m.classification, f"{m.containment:.4f}", f"{m.jaccard:.4f}", f"{m.image_similarity:.4f}", f"{m.title_similarity:.4f}", a.path, b.path, a.size, b.size, a.pages or "", b.pages or "", extra_hint(a, b)])

    with err_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["file", "extension", "error"])
        for fp in fps:
            if fp.error:
                w.writerow([fp.path, fp.ext, fp.error])

    match_map: dict[tuple[int, int], Match] = {}
    for m in matches:
        match_map[tuple(sorted((m.a, m.b)))] = m

    cards = []
    for gi, members in enumerate(groups, 1):
        files_html = []
        for idx in members:
            fp = fps[idx]
            files_html.append(f"""
            <div class='file'>
              <div><b>{esc(Path(fp.path).name)}</b> <span class='pill'>{esc(fp.ext[1:].upper())}</span></div>
              <div class='path'>{esc(fp.path)}</div>
              <div class='meta'>Size: {esc(fmt_size(fp.size))} &nbsp; | &nbsp; Pages: {esc(fp.pages if fp.pages is not None else 'n/a')} &nbsp; | &nbsp; Words: {fp.word_count:,}</div>
            </div>""")
        pair_rows = []
        for x in range(len(members)):
            for y in range(x + 1, len(members)):
                key = tuple(sorted((members[x], members[y])))
                m = match_map.get(key)
                if not m:
                    continue
                a, b = fps[m.a], fps[m.b]
                pair_rows.append(f"""
                  <tr>
                    <td>{esc(m.classification)}</td>
                    <td>{m.containment*100:.1f}%</td>
                    <td>{m.jaccard*100:.1f}%</td>
                    <td>{m.image_similarity*100:.1f}%</td>
                    <td>{m.title_similarity*100:.1f}%</td>
                    <td>{esc(Path(a.path).name)} ↔ {esc(Path(b.path).name)}</td>
                    <td>{esc(extra_hint(a,b))}</td>
                  </tr>""")
        cards.append(f"""
        <section class='group'>
          <h2>Group {gi:04d} <span class='count'>{len(members)} files</span></h2>
          {''.join(files_html)}
          <table>
            <thead><tr><th>Class</th><th>Containment</th><th>Jaccard</th><th>Image</th><th>Title</th><th>Pair</th><th>Difference hint</th></tr></thead>
            <tbody>{''.join(pair_rows)}</tbody>
          </table>
        </section>""")

    error_count = sum(bool(x.error) for x in fps)
    html_doc = f"""<!doctype html>
<html><head><meta charset='utf-8'><title>Book Duplicate Report</title>
<style>
body{{font-family:Segoe UI,Arial,sans-serif;background:#f4f6f8;color:#202124;margin:0}}
header{{background:#1f2937;color:white;padding:24px 30px;position:sticky;top:0;z-index:2}}
main{{max-width:1500px;margin:auto;padding:24px}}
.summary{{display:flex;gap:16px;flex-wrap:wrap;margin-bottom:20px}}
.stat{{background:white;padding:14px 18px;border-radius:10px;box-shadow:0 1px 4px #0002;min-width:160px}}
.group{{background:white;border-radius:12px;margin:0 0 24px;padding:18px;box-shadow:0 1px 5px #0002}}
.file{{padding:9px 10px;border-left:4px solid #64748b;margin:8px 0;background:#f8fafc}}
.path{{font-family:Consolas,monospace;font-size:12px;word-break:break-all;margin-top:3px}}
.meta{{font-size:12px;color:#475569;margin-top:3px}}
.pill{{font-size:11px;background:#e2e8f0;border-radius:999px;padding:2px 7px}}
.count{{font-size:13px;font-weight:normal;color:#64748b}}
table{{border-collapse:collapse;width:100%;margin-top:14px;font-size:13px}}
th,td{{border:1px solid #d7dde5;padding:7px;vertical-align:top;text-align:left}}
th{{background:#eef2f7}}
.note{{background:#fff7d6;border:1px solid #eedb8c;padding:12px;border-radius:8px;margin-bottom:20px}}
</style></head>
<body><header><h1 style='margin:0'>Book / Document Content Duplicate Report</h1><div>Content-aware matching across different sizes, page counts and formats</div></header>
<main>
<div class='summary'>
 <div class='stat'><b>{len(fps):,}</b><br>files scanned</div>
 <div class='stat'><b>{len(matches):,}</b><br>matching pairs</div>
 <div class='stat'><b>{len(groups):,}</b><br>duplicate groups</div>
 <div class='stat'><b>{error_count:,}</b><br>extraction errors</div>
 <div class='stat'><b>{elapsed/60:.1f} min</b><br>elapsed</div>
</div>
<div class='note'><b>No files were deleted or moved.</b> HIGH/PROBABLE/REVIEW are similarity classifications, not deletion instructions. Extra pages or alternate editions should be reviewed before removing anything.</div>
{''.join(cards) if cards else '<section class="group"><h2>No duplicate groups met the configured thresholds.</h2></section>'}
</main></body></html>"""
    html_path.write_text(html_doc, encoding="utf-8")
    return html_path, csv_path, err_path


# ---------------------------------------------------------------------------
# Main program
# ---------------------------------------------------------------------------
def main() -> int:
    started = time.time()
    search_roots = read_path_file("search_paths.txt")
    excludes = read_path_file("exclude_paths.txt")

    print("Search paths:")
    for p in search_roots:
        print("  +", p)
    print("Excluded paths:")
    for p in excludes:
        print("  -", p)

    if not search_roots:
        print("\nERROR: search_paths.txt has no active folders.")
        return 2

    print("\nFinding supported files...")
    files = scan_files(search_roots, excludes)
    print(f"Found {len(files):,} supported files.")
    if not files:
        print("Nothing to scan. Edit search_paths.txt and run again.")
        return 3

    con = db_connect()
    fps: list[Fingerprint] = []
    new_count = 0
    cached_count = 0
    try:
        for n, path in enumerate(files, 1):
            cached = get_cached(con, path)
            if cached:
                fp = cached
                cached_count += 1
                status = "cached"
            else:
                fp = build_fp(path)
                save_cached(con, fp)
                con.commit()
                new_count += 1
                status = "ERROR" if fp.error else "analyzed"
            fps.append(fp)
            print(f"[{n}/{len(files)}] {status}: {path.name}")
            if fp.error:
                print("    ", fp.error)
    finally:
        con.commit()
        con.close()

    print(f"\nFingerprinting complete: {new_count:,} analyzed, {cached_count:,} reused from cache.")
    print("Building candidate pairs...")
    candidates = build_candidates(fps)
    print(f"Candidate pairs: {len(candidates):,}")

    matches: list[Match] = []
    dsu = DSU(len(fps))
    for n, (ia, ib) in enumerate(sorted(candidates), 1):
        m = classify_pair(fps[ia], fps[ib])
        if m:
            m.a, m.b = ia, ib
            matches.append(m)
            dsu.union(ia, ib)
        if n % 5000 == 0:
            print(f"Compared {n:,}/{len(candidates):,} candidate pairs; matches so far: {len(matches):,}")

    roots: dict[int, list[int]] = defaultdict(list)
    matched_indices = set()
    for m in matches:
        matched_indices.add(m.a); matched_indices.add(m.b)
    for i in matched_indices:
        roots[dsu.find(i)].append(i)
    groups = [sorted(v, key=lambda idx: fps[idx].size, reverse=True) for v in roots.values() if len(v) >= 2]
    groups.sort(key=lambda g: (-len(g), -max(fps[i].size for i in g)))

    elapsed = time.time() - started
    html_path, csv_path, err_path = write_reports(fps, matches, groups, elapsed)
    print("\n============================================================")
    print("FINISHED")
    print("============================================================")
    print(f"Files scanned       : {len(fps):,}")
    print(f"Matching pairs      : {len(matches):,}")
    print(f"Duplicate groups    : {len(groups):,}")
    print(f"Extraction errors   : {sum(bool(x.error) for x in fps):,}")
    print(f"HTML report         : {html_path}")
    print(f"CSV pair report     : {csv_path}")
    print(f"CSV errors          : {err_path}")
    print(f"Cache database      : {CACHE_DB}")
    print("No files were deleted or moved.")

    if OPEN_REPORT:
        try:
            webbrowser.open(html_path.as_uri())
        except Exception:
            pass
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
