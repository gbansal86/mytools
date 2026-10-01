#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import deque
from concurrent.futures import ThreadPoolExecutor, as_completed
import csv
import hashlib
import html as html_lib
import json
import os
import re
import shutil
import sqlite3
import sys
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin, urlparse, urldefrag

import requests
from bs4 import BeautifulSoup

import downloadly_worker as legacy

APP_VERSION = "30.0"
PACKAGE_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DATA = PACKAGE_ROOT / "data"

TOPICS = {
    "video-tutorials": {"name": "Video Tutorial", "url": "https://downloadlynet.ir/topics/video-tutorials/"},
    "software": {"name": "Software", "url": "https://downloadlynet.ir/topics/software/"},
    "utility": {"name": "Utility", "url": "https://downloadlynet.ir/topics/utility/"},
    "driver": {"name": "Driver", "url": "https://downloadlynet.ir/topics/driver/"},
    "data-recovery": {"name": "Data Recovery", "url": "https://downloadlynet.ir/topics/data-recovery/"},
    "network-server": {"name": "Network / Server", "url": "https://downloadlynet.ir/topics/network-server/"},
    "operating-system": {"name": "Operating System", "url": "https://downloadlynet.ir/topics/operating-system/"},
    "antivirus-firewall": {"name": "Antivirus / Firewall", "url": "https://downloadlynet.ir/topics/antivirus-firewall/"},
    "engineering-specialized": {"name": "Engineering / Specialized", "url": "https://downloadlynet.ir/topics/engineering-specialized/"},
    "development": {"name": "Development", "url": "https://downloadlynet.ir/topics/%D8%AA%D9%88%D8%B3%D8%B9%D9%87/"},
    "lifestyle": {"name": "Lifestyle", "url": "https://downloadlynet.ir/topics/%D8%B3%D8%A8%DA%A9-%D8%B2%D9%86%D8%AF%DA%AF%DB%8C/"},
    "personal-development": {"name": "Personal Development", "url": "https://downloadlynet.ir/topics/%D8%AA%D9%88%D8%B3%D8%B9%D9%87-%D8%B4%D8%AE%D8%B5%DB%8C/"},
    "converter": {"name": "Converter", "url": "https://downloadlynet.ir/topics/converter/"},
    "graphic": {"name": "Graphic", "url": "https://downloadlynet.ir/topics/graphic/"},
    "programming": {"name": "Programming", "url": "https://downloadlynet.ir/topics/programming/"},
    "audio-video-editors": {"name": "Audio / Video Editors", "url": "https://downloadlynet.ir/topics/audio-video-editors/"},
    "ebook": {"name": "Ebook", "url": "https://downloadlynet.ir/topics/ebook/"},
    "office-productivity": {"name": "Office Productivity", "url": "https://downloadlynet.ir/topics/%D8%A8%D9%87%D8%B1%D9%87-%D9%88%D8%B1%DB%8C-%D8%AF%D8%B1-%D8%A2%D9%81%DB%8C%D8%B3/"},
}

REPORT_HEADERS = [
    "Title", "English Title", "Original Title", "Topic / Category", "Parent Page URL",
    "Published Date", "Thumbnail URL", "Local Thumbnail", "Udemy Source/coupon Link(s)",
    "Google Drive Link(s)", "Rapidgator Link(s)", "Direct File Link(s)",
    "Other File Host Link(s)", "Host(s)", "Saved HTML", "Status", "Notes",
]


@dataclass
class TranslationConfig:
    engine: str = "disabled"  # disabled | argos
    translate_title: bool = True
    translate_description: bool = False
    auto_install_package: bool = False


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _canonical(url: str) -> str:
    return legacy.canonical_parent_key(url)


def _item_id(url: str) -> str:
    return hashlib.sha1(_canonical(url).encode("utf-8", errors="ignore")).hexdigest()[:20]


def _topic_name(slug: str) -> str:
    return TOPICS.get(slug, {}).get("name", slug)


def detect_language_simple(text: str) -> str:
    text = text or ""
    if not text.strip():
        return "en"
    if re.search(r"[\u0600-\u06FF]", text):
        return "fa"
    if re.search(r"[\u0900-\u097F]", text):
        return "hi"
    if re.search(r"[\u0400-\u04FF]", text):
        return "ru"
    if re.search(r"[\u4E00-\u9FFF]", text):
        return "zh"
    return "en"


def translate_text(text: str, config: TranslationConfig) -> str:
    if not text or config.engine == "disabled":
        return text
    if config.engine != "argos":
        return text
    src = detect_language_simple(text)
    if src == "en":
        return text
    try:
        import argostranslate.package
        import argostranslate.translate
        installed = argostranslate.translate.get_installed_languages()
        src_lang = next((x for x in installed if x.code == src), None)
        en_lang = next((x for x in installed if x.code == "en"), None)
        if (src_lang is None or en_lang is None) and config.auto_install_package:
            argostranslate.package.update_package_index()
            packages = argostranslate.package.get_available_packages()
            pkg = next((p for p in packages if p.from_code == src and p.to_code == "en"), None)
            if pkg:
                argostranslate.package.install_from_path(pkg.download())
                installed = argostranslate.translate.get_installed_languages()
                src_lang = next((x for x in installed if x.code == src), None)
                en_lang = next((x for x in installed if x.code == "en"), None)
        if src_lang and en_lang:
            translation = src_lang.get_translation(en_lang)
            return translation.translate(text)
    except Exception:
        return text
    return text


class CatalogDB:
    def __init__(self, path: Path | str):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._local = threading.local()
        self.init_schema()

    def connect(self):
        con = sqlite3.connect(self.path, timeout=30)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA journal_mode=WAL")
        con.execute("PRAGMA synchronous=NORMAL")
        return con

    def init_schema(self):
        with self.connect() as con:
            con.executescript("""
            CREATE TABLE IF NOT EXISTS items(
                item_id TEXT PRIMARY KEY,
                canonical_url TEXT NOT NULL UNIQUE,
                source_url TEXT NOT NULL,
                original_title TEXT NOT NULL DEFAULT '',
                english_title TEXT NOT NULL DEFAULT '',
                published_date TEXT NOT NULL DEFAULT '',
                image_url TEXT NOT NULL DEFAULT '',
                local_image TEXT NOT NULL DEFAULT '',
                html_path TEXT NOT NULL DEFAULT '',
                processing_state TEXT NOT NULL DEFAULT 'pending',
                processing_error TEXT NOT NULL DEFAULT '',
                retry_count INTEGER NOT NULL DEFAULT 0,
                first_discovered TEXT NOT NULL DEFAULT '',
                last_discovered TEXT NOT NULL DEFAULT '',
                last_processed TEXT NOT NULL DEFAULT '',
                manual_status TEXT NOT NULL DEFAULT '',
                manual_notes TEXT NOT NULL DEFAULT '',
                manual_title TEXT NOT NULL DEFAULT '',
                manual_category TEXT NOT NULL DEFAULT '',
                manual_thumbnail TEXT NOT NULL DEFAULT '',
                duplicate_group TEXT NOT NULL DEFAULT '',
                duplicate_role TEXT NOT NULL DEFAULT '',
                duplicate_reason TEXT NOT NULL DEFAULT '',
                duplicate_score REAL NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS item_topics(
                item_id TEXT NOT NULL,
                topic_slug TEXT NOT NULL,
                PRIMARY KEY(item_id, topic_slug)
            );
            CREATE TABLE IF NOT EXISTS links(
                item_id TEXT NOT NULL,
                link_type TEXT NOT NULL DEFAULT '',
                label TEXT NOT NULL DEFAULT '',
                url TEXT NOT NULL,
                host TEXT NOT NULL DEFAULT '',
                PRIMARY KEY(item_id, url)
            );
            CREATE TABLE IF NOT EXISTS topic_progress(
                topic_slug TEXT PRIMARY KEY,
                last_completed_batch INTEGER NOT NULL DEFAULT 0,
                next_batch INTEGER NOT NULL DEFAULT 1,
                total_batches INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT 'new',
                updated_at TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS app_settings(
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL DEFAULT ''
            );
            """)
            # Upgrade existing catalogs in place without touching collected data.
            cols={r[1] for r in con.execute("PRAGMA table_info(items)").fetchall()}
            for name,ddl in (("duplicate_group","TEXT NOT NULL DEFAULT ''"),("duplicate_role","TEXT NOT NULL DEFAULT ''"),("duplicate_reason","TEXT NOT NULL DEFAULT ''"),("duplicate_score","REAL NOT NULL DEFAULT 0")):
                if name not in cols:
                    con.execute(f"ALTER TABLE items ADD COLUMN {name} {ddl}")

    def upsert_discovered(self, row: dict) -> str:
        url = str(row.get("url") or row.get("source_url") or "").strip()
        if not url:
            raise ValueError("Discovered row requires url")
        iid = _item_id(url)
        title = str(row.get("title") or legacy.fallback_title_from_url(url)).strip()
        published = str(row.get("url_date") or row.get("published_date") or "").strip()
        image = str(row.get("image_url") or "").strip()
        topic = str(row.get("topic_slug") or "").strip()
        stamp = _now()
        with self.connect() as con:
            existing = con.execute("SELECT item_id FROM items WHERE canonical_url=?", (_canonical(url),)).fetchone()
            if existing:
                iid = existing["item_id"]
                con.execute("UPDATE items SET last_discovered=? WHERE item_id=?", (stamp, iid))
            else:
                con.execute("""INSERT INTO items(item_id,canonical_url,source_url,original_title,published_date,image_url,first_discovered,last_discovered)
                             VALUES(?,?,?,?,?,?,?,?)""",
                            (iid, _canonical(url), url, title, published, image, stamp, stamp))
            if topic:
                con.execute("INSERT OR IGNORE INTO item_topics(item_id,topic_slug) VALUES(?,?)", (iid, topic))
        return iid

    def get_item(self, item_id: str) -> dict | None:
        with self.connect() as con:
            row = con.execute("SELECT * FROM items WHERE item_id=?", (item_id,)).fetchone()
            return dict(row) if row else None

    def set_manual(self, item_id: str, *, status: str | None = None, notes: str | None = None,
                   title: str | None = None, category: str | None = None, thumbnail: str | None = None):
        fields=[]; vals=[]
        for col, val in (("manual_status",status),("manual_notes",notes),("manual_title",title),("manual_category",category),("manual_thumbnail",thumbnail)):
            if val is not None:
                fields.append(f"{col}=?"); vals.append(val)
        if not fields: return
        vals.append(item_id)
        with self.connect() as con:
            con.execute(f"UPDATE items SET {', '.join(fields)} WHERE item_id=?", vals)

    def mark_state(self, item_id: str, state: str, error: str = "", *, increment_retry: bool = False):
        with self.connect() as con:
            if increment_retry:
                con.execute("UPDATE items SET processing_state=?,processing_error=?,retry_count=retry_count+1,last_processed=? WHERE item_id=?",
                            (state,error,_now(),item_id))
            else:
                con.execute("UPDATE items SET processing_state=?,processing_error=?,last_processed=? WHERE item_id=?",
                            (state,error,_now(),item_id))

    def recover_interrupted_items(self) -> int:
        """Return rows left in `processing` by an interrupted/force-stopped worker to pending."""
        with self.connect() as con:
            cur = con.execute("UPDATE items SET processing_state='pending', processing_error=CASE WHEN processing_error='' THEN 'Interrupted previous run; queued for retry' ELSE processing_error END WHERE processing_state='processing'")
            return int(cur.rowcount or 0)

    def state_counts(self, topics: list[str] | None = None) -> dict:
        params=[]
        where='1=1'
        if topics:
            t_ph=','.join('?'*len(topics)); params.extend(topics)
            where += f" AND EXISTS(SELECT 1 FROM item_topics t WHERE t.item_id=i.item_id AND t.topic_slug IN ({t_ph}))"
        with self.connect() as con:
            rows=con.execute(f"SELECT processing_state, COUNT(*) n FROM items i WHERE {where} GROUP BY processing_state", params).fetchall()
        out={str(r['processing_state']): int(r['n']) for r in rows}
        out['total']=sum(out.values())
        return out

    def update_processed(self, item_id: str, page, links: list, english_title: str = "", local_image: str = ""):
        html_path = str(getattr(page, "saved_html", "") or "")
        if not html_path:
            html_path = str(getattr(page, "panel_saved_html", "") or "")
        with self.connect() as con:
            con.execute("""UPDATE items SET english_title=CASE WHEN english_title='' THEN ? ELSE english_title END,
                           html_path=CASE WHEN html_path='' THEN ? ELSE html_path END,
                           local_image=CASE WHEN local_image='' THEN ? ELSE local_image END,
                           processing_state='completed',processing_error='',last_processed=? WHERE item_id=?""",
                        (english_title, html_path, local_image, _now(), item_id))
            for link in links:
                url = legacy.unwrap_href_li(str(getattr(link, "child_url", "") or "").strip())
                if not url: continue
                host = urlparse(url).netloc.lower().split(":",1)[0]
                con.execute("INSERT OR IGNORE INTO links(item_id,link_type,label,url,host) VALUES(?,?,?,?,?)",
                            (item_id, str(getattr(link,"link_type","") or ""), str(getattr(link,"link_text","") or ""), url, host))

    def pending_items(self, *, topics: list[str] | None = None, include_failed: bool = True) -> list[dict]:
        states = ("pending","blocked","failed","error") if include_failed else ("pending","blocked")
        placeholders = ",".join("?" * len(states))
        params=list(states)
        where=f"i.processing_state IN ({placeholders})"
        if topics:
            t_ph=",".join("?"*len(topics)); params.extend(topics)
            where += f" AND EXISTS(SELECT 1 FROM item_topics t WHERE t.item_id=i.item_id AND t.topic_slug IN ({t_ph}))"
        with self.connect() as con:
            rows=con.execute(f"SELECT i.* FROM items i WHERE {where} ORDER BY i.published_date DESC, i.first_discovered ASC", params).fetchall()
            return [dict(r) for r in rows]

    def all_items(self) -> list[dict]:
        with self.connect() as con:
            rows=con.execute("SELECT * FROM items ORDER BY published_date DESC, original_title COLLATE NOCASE").fetchall()
            return [dict(r) for r in rows]

    def topics_for(self, item_id: str) -> list[str]:
        with self.connect() as con:
            return [r[0] for r in con.execute("SELECT topic_slug FROM item_topics WHERE item_id=? ORDER BY topic_slug", (item_id,)).fetchall()]

    def links_for(self, item_id: str) -> list[dict]:
        with self.connect() as con:
            return [dict(r) for r in con.execute("SELECT * FROM links WHERE item_id=? ORDER BY host,url", (item_id,)).fetchall()]

    def has_url(self, url: str) -> bool:
        with self.connect() as con:
            return con.execute("SELECT 1 FROM items WHERE canonical_url=?", (_canonical(url),)).fetchone() is not None

    def duplicate_info(self, item_id: str) -> dict:
        with self.connect() as con:
            row=con.execute("SELECT duplicate_group,duplicate_role,duplicate_reason,duplicate_score FROM items WHERE item_id=?",(item_id,)).fetchone()
            return dict(row) if row else {"duplicate_group":"","duplicate_role":"","duplicate_reason":"","duplicate_score":0}

    def duplicate_members(self, group_id: str) -> list[dict]:
        if not group_id: return []
        with self.connect() as con:
            return [dict(r) for r in con.execute("SELECT * FROM items WHERE duplicate_group=? ORDER BY CASE duplicate_role WHEN 'primary' THEN 0 ELSE 1 END,published_date DESC",(group_id,)).fetchall()]

    def clear_duplicate_marks(self):
        with self.connect() as con:
            con.execute("UPDATE items SET duplicate_group='',duplicate_role='',duplicate_reason='',duplicate_score=0")

    def duplicate_rows(self) -> list[dict]:
        with self.connect() as con:
            return [dict(r) for r in con.execute("SELECT * FROM items WHERE duplicate_group<>'' ORDER BY duplicate_group,CASE duplicate_role WHEN 'primary' THEN 0 ELSE 1 END,published_date DESC").fetchall()]

    def _export_topic_progress(self):
        try:
            with self.connect() as con:
                rows=[dict(r) for r in con.execute('SELECT * FROM topic_progress ORDER BY topic_slug').fetchall()]
            legacy.atomic_write_json(self.path.parent/'topic_progress.json', {'version':22,'topics':rows,'updated_at':_now()})
        except Exception:
            pass

    def save_topic_progress(self, topic_slug: str, *, last_completed_batch: int, next_batch: int, total_batches: int, status: str):
        with self.connect() as con:
            con.execute("""INSERT INTO topic_progress(topic_slug,last_completed_batch,next_batch,total_batches,status,updated_at)
                         VALUES(?,?,?,?,?,?) ON CONFLICT(topic_slug) DO UPDATE SET
                         last_completed_batch=excluded.last_completed_batch,next_batch=excluded.next_batch,
                         total_batches=excluded.total_batches,status=excluded.status,updated_at=excluded.updated_at""",
                        (topic_slug,int(last_completed_batch),int(next_batch),int(total_batches),status,_now()))
        self._export_topic_progress()

    def topic_progress(self, topic_slug: str) -> dict | None:
        with self.connect() as con:
            row=con.execute("SELECT * FROM topic_progress WHERE topic_slug=?", (topic_slug,)).fetchone()
            return dict(row) if row else None

    def reset_topic_progress(self, topic_slug: str):
        with self.connect() as con:
            con.execute("DELETE FROM topic_progress WHERE topic_slug=?", (topic_slug,))
        self._export_topic_progress()

    def count_items(self) -> int:
        with self.connect() as con:
            return int(con.execute("SELECT COUNT(*) FROM items").fetchone()[0])


def extract_topic_items(html: str, listing_url: str, topic_slug: str) -> list[dict]:
    soup=BeautifulSoup(html or "", "lxml")
    base_host=urlparse(listing_url).netloc.lower().split(":",1)[0] or "downloadlynet.ir"
    found={}; order=[]
    for a in soup.find_all("a", href=True):
        absolute=urljoin(listing_url, str(a.get("href") or ""))
        p=urlparse(urldefrag(absolute)[0])
        clean=f"{p.scheme}://{p.netloc}{p.path}"
        if not legacy.is_downloadly_title_url(clean, base_host=base_host):
            continue
        key=_canonical(clean)
        if key in found:
            continue
        card=a.find_parent(class_=lambda v: v and 'w-grid-item' in (v if isinstance(v,list) else str(v).split()))
        title=" ".join(a.stripped_strings).strip()
        image=""; published=""
        scope=card or a.parent or a
        if not title:
            img=scope.find("img") if hasattr(scope,"find") else None
            if img: title=str(img.get("alt") or "").strip()
        if hasattr(scope,"find"):
            img=scope.find("img")
            if img:
                image=str(img.get("data-src") or img.get("src") or "").strip()
                if image: image=urljoin(listing_url,image)
            tm=scope.find("time")
            if tm:
                dt=str(tm.get("datetime") or "").strip()
                if dt:
                    published=dt[:10]
        d=legacy.parse_downloadly_url_date(clean)
        if not published and d: published=d.isoformat()
        found[key]={"title": title or legacy.fallback_title_from_url(clean), "url":clean,
                    "url_date":published, "image_url":image, "topic_slug":topic_slug}
        order.append(key)
    return [found[k] for k in order]


def extract_generic_ajax_config(html_text: str, listing_url: str) -> dict | None:
    soup=BeautifulSoup(html_text or "", "lxml")
    candidates=[]
    for node in soup.select('.w-grid-json[onclick]'):
        raw=str(node.get('onclick') or '').strip()
        m=re.match(r'^\s*return\s+(\{.*\})\s*;?\s*$', raw, re.S)
        if not m: continue
        try: cfg=json.loads(m.group(1))
        except Exception: continue
        if not isinstance(cfg,dict) or cfg.get('action')!='us_ajax_grid': continue
        grid=node.find_parent(class_=lambda v:v and 'w-grid' in (v if isinstance(v,list) else str(v).split()))
        fragment=str(grid) if grid is not None else html_text
        title_count=len(extract_topic_items(fragment,listing_url,''))
        item_count=len(grid.select('.w-grid-item')) if grid else 0
        load=1 if grid and grid.select_one('.g-loadmore') else 0
        try: pages=max(1,int(cfg.get('max_num_pages') or 1))
        except Exception: pages=1
        candidates.append(((load,title_count,item_count,pages),dict(cfg)))
    if not candidates: return None
    candidates.sort(key=lambda x:x[0], reverse=True)
    cfg=candidates[0][1]
    cfg['ajax_url']=str(cfg.get('ajax_url') or urljoin(listing_url,'wp-admin/admin-ajax.php'))
    try: cfg['max_num_pages']=max(1,int(cfg.get('max_num_pages') or 1))
    except Exception: cfg['max_num_pages']=1
    return cfg


def _retryable_file_error(exc: OSError) -> bool:
    return (
        isinstance(exc, PermissionError)
        or getattr(exc, "winerror", None) in {5, 32, 33}
        or getattr(exc, "errno", None) in {13, 16}
    )


def atomic_replace(tmp: Path | str, target: Path | str, *, retries: int = 20, base_delay: float = 0.05) -> None:
    """Replace *target* with *tmp*, retrying brief Windows sharing violations."""
    tmp, target = Path(tmp), Path(target)
    attempts = max(1, int(retries or 1))
    for attempt in range(attempts):
        try:
            os.replace(tmp, target)
            return
        except OSError as exc:
            if not _retryable_file_error(exc) or attempt + 1 >= attempts:
                raise
            time.sleep(max(0.0, float(base_delay)) * (attempt + 1))


def _unique_temp(target: Path, suffix: str = ".tmp") -> Path:
    return target.parent / f"{target.name}.{os.getpid()}.{threading.get_ident()}.{uuid.uuid4().hex}{suffix}"


def write_discovery_csv(db: CatalogDB, data_dir: Path) -> Path:
    path=Path(data_dir)/"discovery"/"discovered_items.csv"; path.parent.mkdir(parents=True,exist_ok=True)
    headers=["Item ID","Title","Topic","Parent Page URL","Published Date","Thumbnail URL","Processing State","First Discovered","Last Discovered"]
    tmp=_unique_temp(path)
    with tmp.open("w",encoding="utf-8-sig",newline="") as f:
        w=csv.DictWriter(f,fieldnames=headers); w.writeheader()
        for item in db.all_items():
            topics=", ".join(_topic_name(x) for x in db.topics_for(item['item_id']))
            w.writerow({"Item ID":item['item_id'],"Title":item['original_title'],"Topic":topics,"Parent Page URL":item['source_url'],
                        "Published Date":item['published_date'],"Thumbnail URL":item['image_url'],"Processing State":item['processing_state'],
                        "First Discovered":item['first_discovered'],"Last Discovered":item['last_discovered']})
    atomic_replace(tmp,path)
    return path


def _group_links(links: list[dict]) -> dict:
    out={"udemy":[],"gdrive":[],"rapid":[],"direct":[],"other":[],"hosts":[]}
    for l in links:
        url=str(l.get('url') or ''); host=str(l.get('host') or '')
        typ=str(l.get('link_type') or ''); label=str(l.get('label') or '').strip()
        disp=f"{label} | {url}" if label and label.lower() not in {'download','link'} else url
        low=host.lower()
        if 'udemy.com' in low: out['udemy'].append(url)
        elif host in {'drive.google.com','docs.google.com'} or host.endswith('.drive.google.com'): out['gdrive'].append(disp)
        elif 'rapidgator' in low: out['rapid'].append(disp)
        elif typ=='Direct file/archive' or ('downloadly' in low and legacy.file_ext(url) in legacy.FILE_EXTS): out['direct'].append(disp)
        elif typ in {'File host / cloud','External download/mirror'}: out['other'].append(disp)
        if host and host not in out['hosts']: out['hosts'].append(host)
    for k in out: out[k]=list(dict.fromkeys(out[k]))
    return out


def _normalized_duplicate_title(title: str) -> str:
    t=(title or '').lower().replace('–','-').replace('—','-')
    t=re.sub(r'^(udemy|linkedin learning|pluralsight|coursera|skillshare|oreilly|o.?reilly)\s*[-:|]\s*','',t)
    t=re.sub(r'\b20\d{2}(?:[-_. ]?\d+)?\b',' ',t)
    t=re.sub(r'\b(v|ver|version)\s*\d+(?:\.\d+)*\b',' ',t)
    t=re.sub(r'[^a-z0-9]+',' ',t)
    return ' '.join(t.split())

def _logical_links_for(db: CatalogDB, item: dict) -> list[dict]:
    group=str(item.get('duplicate_group') or '')
    if group and str(item.get('duplicate_role') or '')=='primary':
        rows=[]; seen=set()
        for member in db.duplicate_members(group):
            for link in db.links_for(member['item_id']):
                key=str(link.get('url') or '')
                if key and key not in seen:
                    seen.add(key); rows.append(link)
        return rows
    return db.links_for(item['item_id'])

def run_duplicate_checker(db: CatalogDB) -> dict:
    """Conservatively group repeat posts without deleting any source record or HTML.

    Exact shared external links are the strongest signal.  Exact normalized titles are
    also accepted when reasonably specific.  One member is marked primary; all source
    posts remain in SQLite and the archive.
    """
    db.clear_duplicate_marks()
    items=db.all_items()
    parent={i['item_id']:i['item_id'] for i in items}
    reason_by_pair={}
    def find(x):
        while parent[x]!=x:
            parent[x]=parent[parent[x]]; x=parent[x]
        return x
    def union(a,b,reason,score):
        ra,rb=find(a),find(b)
        if ra!=rb: parent[rb]=ra
        reason_by_pair[tuple(sorted((a,b)))]=(reason,score)
    by_title={}
    by_link={}
    for item in items:
        iid=item['item_id']
        norm=_normalized_duplicate_title(item.get('original_title') or '')
        if len(norm)>=12:
            by_title.setdefault(norm,[]).append(iid)
        for link in db.links_for(iid):
            u=str(link.get('url') or '').strip()
            host=str(link.get('host') or '').lower()
            if not u or 'downloadly' in host: continue
            if host in {'www.udemy.com','udemy.com','drive.google.com','docs.google.com'} or 'rapidgator' in host or str(link.get('link_type') or '')=='Official course/source':
                by_link.setdefault(u,[]).append(iid)
    for norm,ids in by_title.items():
        if len(ids)>1:
            base=ids[0]
            for x in ids[1:]: union(base,x,'same normalized title',90)
    for url,ids in by_link.items():
        if len(ids)>1:
            base=ids[0]
            for x in ids[1:]: union(base,x,'shared extracted link',100)
    groups={}
    for iid in parent:
        groups.setdefault(find(iid),[]).append(iid)
    group_count=dup_count=0
    with db.connect() as con:
        for members in groups.values():
            if len(members)<2: continue
            group_count+=1
            rows=[db.get_item(i) for i in members]
            def rank(row):
                return (1 if row.get('processing_state')=='completed' else 0, len(db.links_for(row['item_id'])), 1 if row.get('html_path') else 0, str(row.get('published_date') or ''))
            primary=max(rows,key=rank)['item_id']
            gid='dup_'+hashlib.sha1('|'.join(sorted(members)).encode()).hexdigest()[:12]
            for iid in members:
                reasons=[]; score=0
                for other in members:
                    if other==iid: continue
                    r=reason_by_pair.get(tuple(sorted((iid,other))))
                    if r: reasons.append(r[0]); score=max(score,r[1])
                role='primary' if iid==primary else 'duplicate'
                if role=='duplicate': dup_count+=1
                con.execute("UPDATE items SET duplicate_group=?,duplicate_role=?,duplicate_reason=?,duplicate_score=? WHERE item_id=?",(gid,role,'; '.join(sorted(set(reasons))) or 'duplicate group',score,iid))
    return {'groups':group_count,'duplicates':dup_count,'items':len(items)}

def build_master_rows(db: CatalogDB) -> list[dict]:
    rows=[]
    for item in db.all_items():
        if str(item.get('duplicate_role') or '')=='duplicate':
            continue
        grouped=_group_links(_logical_links_for(db,item))
        topics=", ".join(_topic_name(x) for x in db.topics_for(item['item_id']))
        title=item['manual_title'] or item['english_title'] or item['original_title']
        rows.append({
            "Title":title,"English Title":item['english_title'],"Original Title":item['original_title'],
            "Topic / Category":item['manual_category'] or topics,"Parent Page URL":item['source_url'],
            "Published Date":item['published_date'],"Thumbnail URL":item['manual_thumbnail'] or item['image_url'],
            "Local Thumbnail":item['local_image'],"Udemy Source/coupon Link(s)":"\n".join(grouped['udemy']),
            "Google Drive Link(s)":"\n".join(grouped['gdrive']),"Rapidgator Link(s)":"\n".join(grouped['rapid']),
            "Direct File Link(s)":"\n".join(grouped['direct']),"Other File Host Link(s)":"\n".join(grouped['other']),
            "Host(s)":"\n".join(grouped['hosts']),"Saved HTML":item['html_path'],
            "Status":item['manual_status'],"Notes":item['manual_notes'],
        })
    return rows


def _import_manual_excel_fields(db: CatalogDB, xlsx: Path):
    if not xlsx.exists(): return
    try:
        from openpyxl import load_workbook
        wb=load_workbook(xlsx,read_only=True,data_only=True)
        ws=wb['Master Catalog'] if 'Master Catalog' in wb.sheetnames else wb.active
        headers=[str(c.value or '') for c in next(ws.iter_rows(min_row=1,max_row=1))]
        idx={h:i for i,h in enumerate(headers)}
        if 'Parent Page URL' not in idx: return
        with db.connect() as con:
            byurl={r['canonical_url']:r['item_id'] for r in con.execute('SELECT item_id,canonical_url FROM items')}
        for row in ws.iter_rows(min_row=2,values_only=True):
            url=str(row[idx['Parent Page URL']] or '')
            iid=byurl.get(_canonical(url))
            if not iid: continue
            status=str(row[idx['Status']] or '') if 'Status' in idx else None
            notes=str(row[idx['Notes']] or '') if 'Notes' in idx else None
            db.set_manual(iid,status=status,notes=notes)
    except Exception:
        return


def write_master_reports(db: CatalogDB, data_dir: Path) -> tuple[Path,Path]:
    """Update the user-facing reports without rebuilding existing completed Excel rows.

    Existing rows that already contain a Saved HTML value are treated as historical
    records and are left untouched.  Pending/discovered rows may be enriched as they
    become processed.  Manual Status/Notes and any extra user-added columns/styles are
    preserved.  New items are appended; rows are never deleted or reordered.
    """
    data_dir=Path(data_dir); rep=data_dir/'reports'; rep.mkdir(parents=True,exist_ok=True)
    csv_path=rep/'master_catalog.csv'; xlsx=rep/'master_catalog.xlsx'
    _import_manual_excel_fields(db,xlsx)
    rows=build_master_rows(db)

    # CSV is generated from the database, but replacement tolerates brief Windows locks.
    tmp=_unique_temp(csv_path)
    try:
        with tmp.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=REPORT_HEADERS); w.writeheader(); w.writerows(rows)
        atomic_replace(tmp,csv_path)
    finally:
        try: tmp.unlink(missing_ok=True)
        except OSError: pass

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    if xlsx.exists():
        wb=load_workbook(xlsx)
        ws=wb['Master Catalog'] if 'Master Catalog' in wb.sheetnames else wb.active
        ws.title='Master Catalog'
        headers=[str(c.value or '') for c in ws[1]] if ws.max_row else []
        for h in REPORT_HEADERS:
            if h not in headers:
                headers.append(h); ws.cell(1,len(headers)).value=h
    else:
        wb=Workbook(); ws=wb.active; ws.title='Master Catalog'
        headers=list(REPORT_HEADERS); ws.append(headers)

    idx={h:i+1 for i,h in enumerate(headers)}
    url_col=idx['Parent Page URL']; saved_col=idx['Saved HTML']
    existing={}
    for r in range(2,ws.max_row+1):
        url=str(ws.cell(r,url_col).value or '').strip()
        if url: existing[_canonical(url)]=r

    new_rows=[]
    for row in rows:
        key=_canonical(row['Parent Page URL'])
        r=existing.get(key)
        if r is None:
            r=ws.max_row+1
            for h in REPORT_HEADERS:
                ws.cell(r,idx[h]).value=row.get(h,'')
            existing[key]=r; new_rows.append(r)
        else:
            # Once a record has valid Saved HTML, freeze that Excel row. This protects
            # prior records from future refreshes and preserves all user edits/formatting.
            already_complete=bool(str(ws.cell(r,saved_col).value or '').strip())
            if not already_complete:
                for h in REPORT_HEADERS:
                    if h in {'Status','Notes'}: continue
                    cell=ws.cell(r,idx[h])
                    if cell.value in (None,''):
                        cell.value=row.get(h,'')

    # Style only headers and newly appended rows; existing user formatting is untouched.
    fill=PatternFill('solid',fgColor='1F4E78'); font=Font(bold=True,color='FFFFFF')
    for cell in ws[1]:
        if cell.column <= len(REPORT_HEADERS) or cell.value in REPORT_HEADERS:
            cell.fill=fill; cell.font=font; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    default_widths=dict(zip(REPORT_HEADERS,[38,38,38,28,55,14,52,42,55,55,55,60,55,32,50,18,40]))
    for h,wid in default_widths.items():
        col=idx[h]; letter=get_column_letter(col)
        if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < 12:
            ws.column_dimensions[letter].width=wid
    for r in new_rows:
        for cell in ws[r]: cell.alignment=Alignment(vertical='top',wrap_text=True)

    # Generated duplicate-review sheet. Source posts are never deleted; this sheet may
    # be safely regenerated on every report refresh.
    if 'Duplicates' in wb.sheetnames:
        del wb['Duplicates']
    dws=wb.create_sheet('Duplicates')
    dup_headers=['Duplicate Group','Role','Title','Source URL','Published Date','Topic / Category','Reason','Score','Saved HTML']
    dws.append(dup_headers)
    for item in db.duplicate_rows():
        topics=', '.join(_topic_name(x) for x in db.topics_for(item['item_id']))
        dws.append([item.get('duplicate_group',''),item.get('duplicate_role',''),item.get('original_title',''),item.get('source_url',''),item.get('published_date',''),topics,item.get('duplicate_reason',''),item.get('duplicate_score',0),item.get('html_path','')])
    for cell in dws[1]:
        cell.fill=fill; cell.font=font; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    dws.freeze_panes='A2'; dws.auto_filter.ref=dws.dimensions
    for col,wid in enumerate([22,12,48,62,14,28,32,10,50],1): dws.column_dimensions[get_column_letter(col)].width=wid
    for row in dws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)

    t=_unique_temp(xlsx,'.tmp.xlsx')
    try:
        wb.save(t)
        # Validate before replacing the user's existing workbook.
        with zipfile.ZipFile(t,'r') as zh:
            if zh.testzip() is not None: raise RuntimeError('Generated XLSX failed ZIP integrity validation')
        atomic_replace(t,xlsx)
    finally:
        try: t.unlink(missing_ok=True)
        except OSError: pass
    return csv_path,xlsx


def refresh_generated_outputs(db: CatalogDB, data_dir: Path, *, discovery: bool=True, reports: bool=True, catalog: bool=True) -> list[str]:
    """Best-effort output refresh. A locked report must never terminate the worker."""
    errors=[]
    actions=[]
    if discovery: actions.append(('discovery CSV', lambda: write_discovery_csv(db,data_dir)))
    if reports: actions.append(('master reports', lambda: write_master_reports(db,data_dir)))
    if catalog: actions.append(('HTML catalog', lambda: write_html_catalog(db,data_dir)))
    for label,fn in actions:
        try: fn()
        except Exception as exc:
            errors.append(f'{label}: {exc}')
            try: log(data_dir,f'Output refresh warning ({label}): {exc}')
            except Exception: pass
    return errors


HTML_STYLE = r'''
:root{--bg:#07111e;--panel:#101d2d;--panel2:#132235;--text:#eaf2ff;--muted:#95a7bd;--blue:#2497ff;--border:#263a50;--green:#00c896}
*{box-sizing:border-box}body{margin:0;font-family:Segoe UI,Arial,sans-serif;background:#f5f7fb;color:#172033}.top{height:74px;background:var(--bg);color:white;display:flex;align-items:center;padding:0 26px;gap:28px}.brand{font-weight:800;font-size:22px;letter-spacing:.3px}.search{flex:1;max-width:620px;background:#132235;border:1px solid #28425e;color:white;border-radius:6px;padding:12px 15px}.controls{display:flex;gap:9px;align-items:center}.controls select,.controls button{padding:9px 11px;border:1px solid var(--border);border-radius:5px;background:#142439;color:#dcecff}.app{display:grid;grid-template-columns:220px 1fr 300px;min-height:calc(100vh - 74px)}.side{background:#0b1726;color:#dbe8f7;padding:22px 15px;border-right:1px solid #263a50}.side h3{font-size:12px;color:#7890aa;letter-spacing:1.5px}.cat{padding:9px 10px;border-radius:5px;cursor:pointer;margin:2px 0}.cat.active,.cat:hover{background:#152a40;color:#fff}.main{padding:24px;min-width:0}.summary{display:flex;justify-content:space-between;align-items:center;margin-bottom:18px}.cards{display:grid;grid-template-columns: repeat(3, minmax(0, 1fr));gap:18px}.card{background:#fff;border:1px solid #dbe4ee;border-radius:9px;overflow:hidden;box-shadow:0 2px 8px rgba(20,38,60,.06);display:flex;flex-direction:column;min-width:0}.thumb{height:176px;background:#e9eef5 center/cover no-repeat;position:relative}.badge{position:absolute;top:12px;left:12px;background:#1976d2;color:white;border-radius:4px;padding:5px 8px;font-size:11px}.cardbody{padding:15px;display:flex;flex-direction:column;gap:9px;flex:1}.title{font-size:16px;font-weight:750;line-height:1.25;color:#14233a;text-decoration:none}.meta{font-size:12px;color:#6d7c90}.links{font-size:12px;max-height:120px;overflow:auto;border-top:1px solid #edf1f5;padding-top:8px}.links a{display:block;color:#1675ca;text-decoration:none;margin:4px 0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.actions{margin-top:auto;display:flex;gap:7px;flex-wrap:wrap}.btn{border:1px solid #ccd8e5;background:#fff;border-radius:5px;padding:7px 9px;font-size:12px;cursor:pointer;color:#243950;text-decoration:none}.btn.primary{background:#167bd4;color:white;border-color:#167bd4}.detail{background:#0e1c2d;color:#eaf2ff;border-left:1px solid #263a50;padding:22px 18px;overflow:auto}.detail img{width:100%;border-radius:7px;margin:8px 0 14px}.detail h2{font-size:19px}.detail .field{margin:12px 0}.detail .label{font-size:11px;color:#7f95ae;text-transform:uppercase}.detail a{color:#56b5ff}.pager{display:flex;justify-content:center;gap:7px;margin:22px}.pagebtn{border:1px solid #ccd7e3;background:white;border-radius:4px;padding:7px 10px;cursor:pointer}.modal{display:none;position:fixed;inset:0;background:rgba(2,8,15,.64);align-items:center;justify-content:center;z-index:20}.modalbox{width:min(620px,90vw);background:white;padding:22px;border-radius:10px}.modalbox label{display:block;font-size:12px;font-weight:700;margin-top:10px}.modalbox input,.modalbox textarea{width:100%;padding:9px;border:1px solid #ccd8e5;border-radius:5px}.tableview{display:none;background:white;border:1px solid #dbe4ee;border-radius:8px;overflow:auto}.tableview table{border-collapse:collapse;width:100%}.tableview th,.tableview td{padding:10px;border-bottom:1px solid #edf1f5;text-align:left;font-size:12px}.empty{padding:50px;text-align:center;color:#8190a1}@media(max-width:1200px){.app{grid-template-columns:190px 1fr}.detail{display:none}}@media(max-width:900px){.app{grid-template-columns:1fr}.side{display:none}.cards{grid-template-columns:1fr}}
'''


def _safe_uri(path: str) -> str:
    if not path: return ''
    try:
        p=Path(path)
        if p.exists(): return p.resolve().as_uri()
    except Exception: pass
    return ''


def catalog_payload(db: CatalogDB) -> list[dict]:
    out=[]
    for item in db.all_items():
        links=_logical_links_for(db,item); grouped=_group_links(links)
        topics=[_topic_name(x) for x in db.topics_for(item['item_id'])]
        out.append({
            'id':item['item_id'],'title':item['manual_title'] or item['english_title'] or item['original_title'],
            'original_title':item['original_title'],'category':item['manual_category'] or ', '.join(topics),
            'topics':topics,'url':item['source_url'],'published':item['published_date'],
            'image':item['manual_thumbnail'] or _safe_uri(item['local_image']) or item['image_url'],
            'saved_html':_safe_uri(item['html_path']),'status':item['manual_status'],'notes':item['manual_notes'],
            'links':links,'hosts':grouped['hosts'],
            'duplicate_group':item.get('duplicate_group') or '', 'duplicate_role':item.get('duplicate_role') or '',
            'duplicate_reason':item.get('duplicate_reason') or '',
            'duplicate_sources':[{'title':m.get('original_title',''),'url':m.get('source_url','')} for m in db.duplicate_members(item.get('duplicate_group') or '')]
            if (item.get('duplicate_group') and item.get('duplicate_role')=='primary') else []
        })
    return out


def write_html_catalog(db: CatalogDB, data_dir: Path) -> Path:
    data_dir=Path(data_dir); catdir=data_dir/'catalog'; catdir.mkdir(parents=True,exist_ok=True)
    path=catdir/'catalog.html'; payload=json.dumps(catalog_payload(db),ensure_ascii=False).replace('</','<\\/')
    categories=sorted({_topic_name(s) for i in db.all_items() for s in db.topics_for(i['item_id'])})
    cats=''.join(f'<div class="cat" data-cat="{html_lib.escape(x)}">{html_lib.escape(x)}</div>' for x in categories)
    doc=f'''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Downloadly Catalog</title><style>{HTML_STYLE}</style></head><body>
<div class="top"><div class="brand">DOWNLOADLY CATALOG</div><input id="q" class="search" placeholder="Search courses, categories, providers..."><div class="controls"><select id="sort"><option value="new">Newest</option><option value="old">Oldest</option><option value="az">A-Z</option></select><select id="pageSize"><option value="25">25</option><option value="50">50</option><option value="100">100</option><option value="200">200</option></select><button id="cardBtn">▦ Cards</button><button id="tableBtn">☰ Table</button><label style="font-size:12px"><input id="showDup" type="checkbox"> Show duplicates</label><button id="exportEdits">Export edits</button><button id="importEdits">Import edits</button><input id="editsFile" type="file" accept="application/json" hidden></div></div>
<div class="app"><aside class="side"><h3>CATEGORIES</h3><div class="cat active" data-cat="">All</div>{cats}<h3>FILTER</h3><div class="cat" data-status="">All Status</div><div class="cat" data-status="has">Has Status</div><div class="cat" data-status="blank">Blank Status</div></aside><main class="main"><div class="summary"><div><b id="count">0</b> items</div><div id="range"></div></div><div id="cards" class="cards"></div><div id="table" class="tableview"></div><div id="pager" class="pager"></div></main><aside id="detail" class="detail"><h2>Select an item</h2><p>Click a card to review its details.</p></aside></div>
<div id="modal" class="modal"><div class="modalbox"><h2>Edit catalog display</h2><input id="editId" type="hidden"><label>Title</label><input id="editTitle"><label>Category</label><input id="editCategory"><label>Status</label><input id="editStatus"><label>Notes</label><textarea id="editNotes" rows="4"></textarea><label>Thumbnail URL</label><input id="editThumb"><div style="margin-top:16px;display:flex;gap:8px"><button class="btn primary" id="saveEdit">Save</button><button class="btn" id="cancelEdit">Cancel</button></div></div></div>
<script>const RAW={payload}; const KEY='downloadly_catalog_edits_v23'; let edits=JSON.parse(localStorage.getItem(KEY)||'{{}}'); let page=1,cat='',statusFilter='',view='cards'; const $=id=>document.getElementById(id); function merged(x){{let e=edits[x.id]||{{}};return Object.assign({{}},x,{{title:e.title||x.title,category:e.category||x.category,status:e.status!==undefined?e.status:x.status,notes:e.notes!==undefined?e.notes:x.notes,image:e.image||x.image}})}} function filtered(){{let q=$('q').value.trim().toLowerCase();let showDup=$('showDup').checked;let arr=RAW.map(merged).filter(x=>(showDup||x.duplicate_role!=='duplicate')&&(!cat||x.topics.includes(cat)||x.category.includes(cat))&&(!q||[x.title,x.original_title,x.category,x.url,(x.hosts||[]).join(' ')].join(' ').toLowerCase().includes(q))&&(statusFilter===''||(statusFilter==='has'?!!x.status:!x.status)));let s=$('sort').value;if(s==='az')arr.sort((a,b)=>a.title.localeCompare(b.title));else if(s==='old')arr.sort((a,b)=>a.published.localeCompare(b.published));else arr.sort((a,b)=>b.published.localeCompare(a.published));return arr}} function linkHtml(x){{return (x.links||[]).map(l=>`<a href="${{esc(l.url)}}" target="_blank">${{esc(l.label||l.host||l.url)}}</a>`).join('')}} function esc(s){{return String(s||'').replace(/[&<>"']/g,m=>({{'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}}[m]))}} function render(){{let arr=filtered(),size=+$('pageSize').value||25,max=Math.max(1,Math.ceil(arr.length/size));if(page>max)page=max;let start=(page-1)*size,show=arr.slice(start,start+size);$('count').textContent=arr.length;$('range').textContent=arr.length?`${{start+1}}-${{Math.min(start+size,arr.length)}} of ${{arr.length}}`:'';$('cards').innerHTML=show.map(x=>`<article class="card" onclick="selectItem('${{x.id}}')"><div class="thumb" style="background-image:url('${{esc(x.image)}}')"><span class="badge">${{esc((x.topics&&x.topics[0])||x.category||'Item')}}</span></div><div class="cardbody"><a class="title" href="${{esc(x.url)}}" target="_blank" onclick="event.stopPropagation()">${{esc(x.title)}}</a><div class="meta">${{esc(x.category)}} · ${{esc(x.published)}}</div>${{x.status?`<div class="meta"><b>Status:</b> ${{esc(x.status)}}</div>`:''}}<div class="links">${{linkHtml(x)}}</div><div class="actions"><a class="btn primary" href="${{esc(x.url)}}" target="_blank" onclick="event.stopPropagation()">Open Source Page</a>${{x.saved_html?`<a class="btn" href="${{esc(x.saved_html)}}" target="_blank" onclick="event.stopPropagation()">Open Saved HTML</a>`:''}}<button class="btn" onclick="event.stopPropagation();editItem('${{x.id}}')">Edit</button></div></div></article>`).join('')||'<div class="empty">No matching items</div>';$('table').innerHTML='<table><thead><tr><th>Title</th><th>Category</th><th>Date</th><th>Status</th><th>Source</th></tr></thead><tbody>'+show.map(x=>`<tr><td><a href="${{esc(x.url)}}" target="_blank">${{esc(x.title)}}</a></td><td>${{esc(x.category)}}</td><td>${{esc(x.published)}}</td><td>${{esc(x.status)}}</td><td><button class="btn" onclick="editItem('${{x.id}}')">Edit</button></td></tr>`).join('')+'</tbody></table>';$('pager').innerHTML='';for(let p=Math.max(1,page-3);p<=Math.min(max,page+3);p++){{let b=document.createElement('button');b.className='pagebtn';b.textContent=p;if(p===page)b.style.background='#167bd4',b.style.color='white';b.onclick=()=>{{page=p;render()}};$('pager').appendChild(b)}}$('cards').style.display=view==='cards'?'grid':'none';$('table').style.display=view==='table'?'block':'none'}} function selectItem(id){{let x=merged(RAW.find(z=>z.id===id));$('detail').innerHTML=`${{x.image?`<img src="${{esc(x.image)}}">`:''}}<h2>${{esc(x.title)}}</h2><div class="field"><div class="label">Category</div>${{esc(x.category)}}</div><div class="field"><div class="label">Published</div>${{esc(x.published)}}</div><div class="field"><div class="label">Status</div>${{esc(x.status||'')}}</div><div class="field"><div class="label">Notes</div>${{esc(x.notes||'')}}</div><div class="field"><div class="label">Links</div>${{linkHtml(x)}}</div>${{x.duplicate_sources&&x.duplicate_sources.length>1?`<div class="field"><div class="label">Duplicate source posts</div>${{x.duplicate_sources.map(s=>`<a href="${{esc(s.url)}}" target="_blank">${{esc(s.title||s.url)}}</a><br>`).join('')}}</div>`:''}}<div class="field"><a href="${{esc(x.url)}}" target="_blank">Open Source Page</a>${{x.saved_html?`<br><a href="${{esc(x.saved_html)}}" target="_blank">Open Saved HTML</a>`:''}}</div><div class="field"><button class="btn" onclick="editItem('${{x.id}}')">Edit</button></div>`}} function editItem(id){{let x=merged(RAW.find(z=>z.id===id));$('editId').value=id;$('editTitle').value=x.title;$('editCategory').value=x.category;$('editStatus').value=x.status||'';$('editNotes').value=x.notes||'';$('editThumb').value=x.image||'';$('modal').style.display='flex'}} $('saveEdit').onclick=()=>{{let id=$('editId').value;edits[id]={{title:$('editTitle').value,category:$('editCategory').value,status:$('editStatus').value,notes:$('editNotes').value,image:$('editThumb').value}};localStorage.setItem(KEY,JSON.stringify(edits));$('modal').style.display='none';render();selectItem(id)}};$('cancelEdit').onclick=()=>$('modal').style.display='none';$('q').oninput=()=>{{page=1;render()}};$('sort').onchange=render;$('pageSize').onchange=()=>{{page=1;render()}};$('cardBtn').onclick=()=>{{view='cards';render()}};$('tableBtn').onclick=()=>{{view='table';render()}};$('showDup').onchange=()=>{{page=1;render()}};document.querySelectorAll('.cat[data-cat]').forEach(n=>n.onclick=()=>{{document.querySelectorAll('.cat[data-cat]').forEach(x=>x.classList.remove('active'));n.classList.add('active');cat=n.dataset.cat;page=1;render()}});document.querySelectorAll('.cat[data-status]').forEach(n=>n.onclick=()=>{{statusFilter=n.dataset.status;page=1;render()}});$('exportEdits').onclick=()=>{{let blob=new Blob([JSON.stringify(edits,null,2)],{{type:'application/json'}}),a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download='catalog_edits.json';a.click();URL.revokeObjectURL(a.href)}};$('importEdits').onclick=()=>$('editsFile').click();$('editsFile').onchange=e=>{{let f=e.target.files[0];if(!f)return;let r=new FileReader();r.onload=()=>{{try{{edits=Object.assign(edits,JSON.parse(r.result));localStorage.setItem(KEY,JSON.stringify(edits));render()}}catch(_e){{alert('Invalid edits JSON')}}}};r.readAsText(f)}};render();</script></body></html>'''
    tmp=_unique_temp(path); tmp.write_text(doc,encoding='utf-8'); atomic_replace(tmp,path); return path


def inside_allowed_hours(now: datetime, start: str, end: str) -> bool:
    if not start or not end: return True
    try:
        sh,sm=map(int,start.split(':')); eh,em=map(int,end.split(':'))
    except Exception: return True
    cur=now.hour*60+now.minute; a=sh*60+sm; b=eh*60+em
    if a==b: return True
    return a <= cur < b if a < b else (cur >= a or cur < b)


class Control:
    def __init__(self, data_dir: Path, allowed_from: str='', allowed_to: str=''):
        self.state_dir=Path(data_dir)/'state'; self.state_dir.mkdir(parents=True,exist_ok=True)
        self.pause_file=self.state_dir/'pause.requested'; self.stop_file=self.state_dir/'stop.requested'
        self.allowed_from=allowed_from; self.allowed_to=allowed_to
    def clear_stale(self):
        for p in (self.pause_file,self.stop_file):
            try:p.unlink(missing_ok=True)
            except Exception:pass
    def stop(self): return self.stop_file.exists()
    def wait_if_paused_or_scheduled(self, progress_cb=None):
        while self.pause_file.exists() or not inside_allowed_hours(datetime.now(),self.allowed_from,self.allowed_to):
            if self.stop(): raise KeyboardInterrupt
            if progress_cb: progress_cb('paused','Paused by user' if self.pause_file.exists() else 'Paused outside scheduled hours')
            time.sleep(2)


def write_progress(data_dir: Path, payload: dict):
    p=Path(data_dir)/'state'/'run_progress.json'; p.parent.mkdir(parents=True,exist_ok=True)
    legacy.atomic_write_json(p,dict(payload,updated_at=_now(),version=28))


def log(data_dir: Path, msg: str):
    p=Path(data_dir)/'logs'/'run.log'; p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('a',encoding='utf-8') as f:f.write(f"[{_now()}] {msg}\n")


def _save_listing_debug(data_dir: Path, topic_slug: str, batch: int, html: str, blocked: bool=False):
    root=Path(data_dir)/'listing_archive'/topic_slug; root.mkdir(parents=True,exist_ok=True)
    suffix='blocked' if blocked else 'listing'; p=root/f'batch_{batch:05d}_{suffix}.html'
    p.write_text(html or '',encoding='utf-8',errors='replace')


def discover_topic(db: CatalogDB, data_dir: Path, topic_slug: str, *, timeout: int, max_batches: int,
                   block_wait: int, increasing_wait: bool, max_wait: int, control: Control,
                   refresh: bool=False, report_every: int=10, report_seconds: int=60, new_only: bool=False, max_block_retries: int=0):
    info=TOPICS[topic_slug]; url=info['url']
    if refresh: db.reset_topic_progress(topic_slug)
    prog=db.topic_progress(topic_slug)
    preserve_full_progress=bool(new_only and prog)
    current_batch=1
    if prog and prog['status']=='complete' and not refresh and not new_only:
        log(data_dir,f"Discovery already complete for {info['name']}; skipping")
        return
    start=1 if new_only else (max(1,int(prog['next_batch'])) if prog else 1)
    attempt=0
    since_report=0; last_report=time.time()
    while True:
        control.wait_if_paused_or_scheduled(lambda s,m:write_progress(data_dir,{'status':s,'phase':'discovery','topic':info['name'],'message':m}))
        if control.stop(): raise KeyboardInterrupt
        session=legacy.get_http_session()
        try:
            resp=session.get(url,timeout=timeout,allow_redirects=True); html=legacy.response_html_text(resp)
            if resp.status_code>=400 or legacy.is_probable_block_page(html,resp.status_code):
                _save_listing_debug(data_dir,topic_slug,1,html,True)
                raise legacy.DiscoveryBlockedError(url,resp.status_code,'topic challenge/block page')
            cfg=extract_generic_ajax_config(html,str(resp.url or url)); total=int((cfg or {}).get('max_num_pages') or 1)
            if max_batches>0: total=min(total,max_batches)
            first_items=extract_topic_items(html,str(resp.url or url),topic_slug)
            if start<=1:
                first_new=sum(1 for row in first_items if not db.has_url(row.get('url','')))
                for row in first_items: db.upsert_discovered(row)
                if new_only and first_new==0:
                    if not preserve_full_progress: db.save_topic_progress(topic_slug,last_completed_batch=1,next_batch=1,total_batches=total,status='incremental_complete')
                    log(data_dir,f"SCHEDULED DISCOVERY {info['name']} | batch 1 contains no new items; stopping incremental scan")
                    refresh_generated_outputs(db,data_dir,discovery=True,reports=False,catalog=False)
                    return
                if not preserve_full_progress:
                    db.save_topic_progress(topic_slug,last_completed_batch=1,next_batch=2,total_batches=total,status=('incremental_complete' if new_only and total<=1 else ('in_progress' if total>1 else 'complete')))
                elif not new_only:
                    db.save_topic_progress(topic_slug,last_completed_batch=1,next_batch=2,total_batches=total,status='in_progress' if total>1 else 'complete')
                refresh_generated_outputs(db,data_dir,discovery=True,reports=False,catalog=False)
                since_report+=1
                if since_report>=max(1,report_every) or time.time()-last_report>=max(5,report_seconds):
                    refresh_generated_outputs(db,data_dir,discovery=False,reports=True,catalog=True); since_report=0; last_report=time.time()
                _save_listing_debug(data_dir,topic_slug,1,html)
                start=2
            if total<=1:
                if not preserve_full_progress: db.save_topic_progress(topic_slug,last_completed_batch=1,next_batch=(1 if new_only else 2),total_batches=1,status=('incremental_complete' if new_only else 'complete'))
                return
            if not cfg:
                raise RuntimeError('Load More AJAX configuration not found for selected topic')
            for batch in range(max(2,start),total+1):
                current_batch=batch
                control.wait_if_paused_or_scheduled(lambda s,m:write_progress(data_dir,{'status':s,'phase':'discovery','topic':info['name'],'batch':batch,'total_batches':total,'message':m}))
                if control.stop(): raise KeyboardInterrupt
                discovered_before=db.count_items()
                remaining=max(0,total-batch+1)
                write_progress(data_dir,{'status':'running','phase':'discovery','topic':info['name'],'batch':batch,'total_batches':total,'discovered':discovered_before,'remaining_batches':remaining,'message':f'Discovering batch {batch}/{total} ({remaining} including current)'})
                log(data_dir,f"DISCOVERY {info['name']} | batch {batch}/{total} | remaining={remaining} | discovered={discovered_before}")
                r=legacy._fetch_impreza_ajax_page(session,cfg,batch,timeout,str(resp.url or url)); text=legacy.response_html_text(r)
                if r.status_code>=400 or legacy.is_probable_block_page(text,r.status_code):
                    _save_listing_debug(data_dir,topic_slug,batch,text,True)
                    raise legacy.DiscoveryBlockedError(f"{cfg.get('ajax_url')} batch {batch}",r.status_code,'Load More challenge/block page')
                rows=extract_topic_items(text,url,topic_slug)
                new_count=sum(1 for row in rows if not db.has_url(row.get('url','')))
                for row in rows: db.upsert_discovered(row)
                discovered_after=db.count_items()
                log(data_dir,f"DISCOVERY {info['name']} | completed {batch}/{total} | found={len(rows)} | new={new_count} | new_total={discovered_after} | remaining={max(0,total-batch)}")
                _save_listing_debug(data_dir,topic_slug,batch,text)
                if new_only and new_count==0:
                    if not preserve_full_progress: db.save_topic_progress(topic_slug,last_completed_batch=batch,next_batch=1,total_batches=total,status='incremental_complete')
                    refresh_generated_outputs(db,data_dir,discovery=True,reports=False,catalog=False)
                    log(data_dir,f"SCHEDULED DISCOVERY {info['name']} | reached known-only batch {batch}; incremental scan complete")
                    return
                if not preserve_full_progress:
                    db.save_topic_progress(topic_slug,last_completed_batch=batch,next_batch=batch+1,total_batches=total,status=('incremental_complete' if new_only and batch>=total else ('in_progress' if batch<total else 'complete')))
                elif not new_only:
                    db.save_topic_progress(topic_slug,last_completed_batch=batch,next_batch=batch+1,total_batches=total,status='in_progress' if batch<total else 'complete')
                refresh_generated_outputs(db,data_dir,discovery=True,reports=False,catalog=False)
                since_report+=1
                if since_report>=max(1,report_every) or time.time()-last_report>=max(5,report_seconds):
                    refresh_generated_outputs(db,data_dir,discovery=False,reports=True,catalog=True); since_report=0; last_report=time.time()
            return
        except legacy.DiscoveryBlockedError as exc:
            attempt+=1
            if max_block_retries>0 and attempt>=max_block_retries:
                log(data_dir,f"Blocked retry limit reached for {info['name']} discovery; deferring until next run")
                write_progress(data_dir,{'status':'paused_blocked','phase':'discovery','topic':info['name'],'batch':current_batch,'discovered':db.count_items(),'message':'Block retry limit reached; deferred to next scheduled/manual run'})
                refresh_generated_outputs(db,data_dir)
                return
            wait=min(max_wait, block_wait*(attempt if increasing_wait else 1))
            progress=db.topic_progress(topic_slug); retry_batch=(current_batch if new_only else (int(progress['next_batch']) if progress else start))
            log(data_dir,f"Blocked during {info['name']} discovery at batch {retry_batch}; clearing app cookies/session and waiting {wait}s")
            write_progress(data_dir,{'status':'paused_blocked','phase':'discovery','topic':info['name'],'batch':retry_batch,'discovered':db.count_items(),'wait_seconds':wait,'message':str(exc)})
            refresh_generated_outputs(db,data_dir)
            legacy.clear_downloadly_cookies(session); legacy.reset_http_session()
            for _ in range(max(1,wait)):
                if control.stop(): raise KeyboardInterrupt
                control.wait_if_paused_or_scheduled()
                time.sleep(1)
            prog=db.topic_progress(topic_slug); start=(current_batch if new_only else (int(prog['next_batch']) if prog else start))
            continue


def _download_thumbnail(url: str, item_id: str, data_dir: Path, timeout: int=30) -> str:
    if not url: return ''
    try:
        r=requests.get(url,timeout=timeout,stream=True,headers={'User-Agent':'Mozilla/5.0'})
        if r.status_code>=400:return ''
        ctype=str(r.headers.get('content-type') or '').lower(); ext='.jpg'
        if 'png' in ctype:ext='.png'
        elif 'webp' in ctype:ext='.webp'
        elif 'gif' in ctype:ext='.gif'
        p=Path(data_dir)/'images'/f'{item_id}{ext}'; p.parent.mkdir(parents=True,exist_ok=True)
        with p.open('wb') as f:
            for chunk in r.iter_content(65536):
                if chunk:f.write(chunk)
        return str(p)
    except Exception:return ''


def process_queue(db: CatalogDB, data_dir: Path, topics: list[str], *, timeout: int, workers: int,
                  block_wait: int, increasing_wait: bool, max_wait: int, control: Control,
                  download_images: bool, translation: TranslationConfig, report_every: int=100, report_seconds: int=300,
                  retry_failed: bool=True, max_block_retries: int=0):
    """Process the durable queue with bounded parallel fetches and serialized DB writes.

    V28 accepted a ``workers`` option but still processed only ``pending[0]`` in a
    single-thread loop.  It also reloaded the entire pending queue after every item.
    V29 loads the queue once, fetches up to ``workers`` course pages concurrently,
    and writes results back to SQLite on the coordinator thread.  The HTTP helper
    already uses thread-local requests sessions, so cookies/sockets are not shared
    unsafely between fetch workers.
    """
    archive=Path(data_dir)/'html_archive'; archive.mkdir(parents=True,exist_ok=True)
    worker_count=max(1,min(int(workers or 1),100))
    last_flush=time.time(); since=0; failed_this_run=set(); block_attempts_this_run={}
    # Keep expensive browser recovery bounded even when HTTP worker count is high.
    selenium_fallback_gate=threading.BoundedSemaphore(2)
    initial=[x for x in db.pending_items(topics=topics, include_failed=retry_failed)]
    queue=deque(initial)
    counts=db.state_counts(topics); total_selected=int(counts.get('total',0)); completed_count=int(counts.get('completed',0))
    log(data_dir,f"PROCESS QUEUE | total={total_selected} | completed={completed_count} | queued={len(queue)} | workers={worker_count} | report_every={max(1,report_every)} | report_seconds={max(5,report_seconds)}")

    def fetch_one(item):
        try:
            page,links,reused=legacy.fetch_archive_extract_one(
                item['source_url'],archive_root=archive,mode='offline',timeout=timeout,
                wait_seconds=2.0,headless_browser=True,refresh=False,preferred_title=item['original_title'])
            if page.fetch_status=='blocked':
                log(data_dir,f"HEADLESS FALLBACK | HTTP block detected | {item['original_title']}")
                try:
                    with selenium_fallback_gate:
                        page,links,reused=legacy.fetch_archive_extract_one(
                            item['source_url'],archive_root=archive,mode='browser',timeout=timeout,
                            wait_seconds=2.0,headless_browser=True,refresh=True,preferred_title=item['original_title'])
                    if page.fetch_status=='fetched':
                        log(data_dir,f"HEADLESS FALLBACK SUCCESS | links={len(links)} | {item['original_title']}")
                    elif page.fetch_status=='blocked':
                        raise legacy.DiscoveryBlockedError(item['source_url'],int(page.http_status or 200),'headless browser also received challenge/block page')
                except legacy.DiscoveryBlockedError:
                    raise
                except Exception as browser_exc:
                    log(data_dir,f"HEADLESS FALLBACK ERROR | {item['source_url']} | {browser_exc}")
                    raise legacy.DiscoveryBlockedError(item['source_url'],int(page.http_status or 200),'headless browser fallback failed') from browser_exc
            if page.fetch_status in {'error','http_error'}:
                return ('failed',item,page,links,reused,page.error or page.fetch_status)
            return ('ok',item,page,links,reused,'')
        except legacy.DiscoveryBlockedError as exc:
            return ('blocked',item,None,[],False,str(exc))
        except Exception as exc:
            return ('failed',item,None,[],False,str(exc))

    with ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix='course-fetch') as executor:
        while queue:
            control.wait_if_paused_or_scheduled(lambda st,msg:write_progress(data_dir,{'status':st,'phase':'processing','message':msg}))
            if control.stop(): raise KeyboardInterrupt

            wave=[]
            while queue and len(wave)<worker_count:
                item=queue.popleft(); iid=item['item_id']
                if iid in failed_this_run: continue
                db.mark_state(iid,'processing')
                ordinal=min(total_selected,completed_count+len(wave)+1) if total_selected else 0
                log(data_dir,f"PROCESS START {ordinal}/{total_selected} | queued={len(queue)+len(wave)+1} | worker-slot={len(wave)+1}/{worker_count} | {item['original_title']}")
                wave.append((executor.submit(fetch_one,item),item))

            if not wave:
                break

            blocked_for_retry=[]
            for future in as_completed([f for f,_ in wave]):
                if control.stop(): raise KeyboardInterrupt
                status,item,page,links,reused,error=future.result(); iid=item['item_id']
                if status=='ok':
                    english=''
                    if translation.engine!='disabled' and translation.translate_title:
                        english=translate_text(item['original_title'],translation)
                        if english==item['original_title']: english=''
                    local=''
                    if download_images and item['image_url'] and not item['local_image']:
                        local=_download_thumbnail(item['image_url'],iid,data_dir,timeout)
                    db.update_processed(iid,page,links,english_title=english,local_image=local)
                    completed_count+=1; result='reused' if reused else 'fetched'; since+=1
                    log(data_dir,f"PROCESS {completed_count}/{total_selected} | {result} | links={len(links)} | pending={len(queue)} | queued={len(queue)} | {item['original_title']}")
                elif status=='blocked':
                    db.mark_state(iid,'blocked',error,increment_retry=True)
                    block_attempts_this_run[iid]=block_attempts_this_run.get(iid,0)+1
                    attempt_this_run=block_attempts_this_run[iid]
                    if max_block_retries>0 and attempt_this_run>=max_block_retries:
                        failed_this_run.add(iid)
                        log(data_dir,f"PROCESS BLOCKED retry limit reached | deferred to next run | {item['original_title']}")
                    else:
                        blocked_for_retry.append(item)
                        log(data_dir,f"PROCESS BLOCKED | retry={attempt_this_run} | {item['original_title']}")
                else:
                    db.mark_state(iid,'failed',error,increment_retry=True); failed_this_run.add(iid); since+=1
                    log(data_dir,f"PROCESS ERROR | {item['source_url']} | {error}")

                if since>=max(1,report_every) or time.time()-last_flush>=max(5,report_seconds):
                    refresh_generated_outputs(db,data_dir)
                    since=0; last_flush=time.time()

            counts=db.state_counts(topics)
            remaining=int(counts.get('pending',0)+counts.get('blocked',0)+counts.get('failed',0)+counts.get('error',0))
            write_progress(data_dir,{'status':'running','phase':'processing','pending':remaining,
                                    'completed':int(counts.get('completed',0)),
                                    'failed':int(counts.get('failed',0)+counts.get('error',0)),
                                    'blocked':int(counts.get('blocked',0)),'items':int(counts.get('total',0)),
                                    'workers':worker_count,
                                    'message':f"Processed {int(counts.get('completed',0))}/{int(counts.get('total',0))}; {remaining} unfinished; {worker_count} workers"})

            if blocked_for_retry:
                # One shared back-off per blocked wave is much faster than waiting once
                # for every blocked URL, while still respecting the site's block signal.
                max_attempt=max(block_attempts_this_run.get(x['item_id'],1) for x in blocked_for_retry)
                wait=min(max_wait,block_wait*(max_attempt if increasing_wait else 1))
                refresh_generated_outputs(db,data_dir)
                log(data_dir,f"PROCESS BLOCK WAVE | blocked={len(blocked_for_retry)} | wait={wait}s | workers={worker_count}")
                write_progress(data_dir,{'status':'paused_blocked','phase':'processing','wait_seconds':wait,
                                        'pending':remaining,'completed':int(counts.get('completed',0)),
                                        'items':int(counts.get('total',0)),
                                        'message':f'{len(blocked_for_retry)} request(s) blocked; rotating sessions and retrying after {wait}s'})
                try: legacy.clear_downloadly_cookies(legacy.get_http_session())
                except Exception: pass
                legacy.reset_http_session()
                for _ in range(max(1,wait)):
                    if control.stop(): raise KeyboardInterrupt
                    control.wait_if_paused_or_scheduled(); time.sleep(1)
                for item in reversed(blocked_for_retry):
                    queue.appendleft(item)

    refresh_generated_outputs(db,data_dir)


def migrate_existing(db: CatalogDB, data_dir: Path) -> dict:
    data_dir=Path(data_dir); added=0; archives=0
    # Previous discovery state
    for p in (data_dir/'state'/'discovered_courses.json', data_dir/'state'/'discovery_history.json'):
        if p.exists():
            try:
                rows=json.loads(p.read_text(encoding='utf-8-sig',errors='replace'))
                if isinstance(rows,list):
                    for r in rows:
                        if r.get('url'):
                            r=dict(r); r.setdefault('topic_slug','video-tutorials'); db.upsert_discovered(r); added+=1
            except Exception: pass
    # V15-V21 interrupted full-discovery checkpoint (may contain thousands of rows).
    cp = data_dir/'state'/'discovery_checkpoint.json'
    if cp.exists():
        try:
            payload=json.loads(cp.read_text(encoding='utf-8-sig',errors='replace'))
            rows=payload.get('rows') if isinstance(payload,dict) else []
            if isinstance(rows,list):
                for r in rows:
                    if isinstance(r,dict) and r.get('url'):
                        r=dict(r); r.setdefault('topic_slug','video-tutorials'); db.upsert_discovered(r); added+=1
                # Do not carry the old listing-batch cursor forward: V19-V21 scanned the main-site grid,
                # while V23 discovers the selected topic URL directly. The saved URLs are reusable; the old batch
                # number is not comparable to the topic grid's batch number.
        except Exception:
            pass

    # Previous report
    old=data_dir/'reports'/'course_report.csv'
    if old.exists():
        try:
            with old.open(encoding='utf-8-sig',newline='') as f:
                for r in csv.DictReader(f):
                    url=str(r.get('Parent Course Page') or '')
                    if not url:continue
                    iid=db.upsert_discovered({'title':r.get('Course / Post Title') or '', 'url':url,'url_date':r.get('Published Date') or '', 'image_url':'','topic_slug':'video-tutorials'})
                    if r.get('Status'): db.set_manual(iid,status='')  # V21 Status was internal, not user's manual Status.
                    # import links from grouped columns
                    for col,typ in [('Udemy Source/coupon Link(s)','Official course/source'),('Google Drive Link(s)','File host / cloud'),('Rapidgator Link(s)','File host / cloud'),('Direct File Link(s)','Direct file/archive'),('Other File Host Link(s)','External download/mirror')]:
                        for line in str(r.get(col) or '').splitlines():
                            u=line.split(' | ')[-1].strip()
                            if u.startswith('http'):
                                with db.connect() as con: con.execute('INSERT OR IGNORE INTO links(item_id,link_type,label,url,host) VALUES(?,?,?,?,?)',(iid,typ,line.split(' | ')[0] if ' | ' in line else '',u,urlparse(u).netloc.lower()))
        except Exception: pass
    # Existing HTML archive - authoritative cache
    archive=data_dir/'html_archive'
    if archive.exists():
        for url,folder in legacy.scan_archive(archive):
            try:
                page,links=legacy.extract_course_from_archive(folder)
                if page.fetch_status=='blocked' or not page.saved_html: continue
                iid=db.upsert_discovered({'title':page.parent_title,'url':url,'url_date':page.url_date,'image_url':'','topic_slug':'video-tutorials'})
                db.update_processed(iid,page,links); archives+=1
            except Exception: continue
    return {'discovered_imports':added,'archives_imported':archives,'total_items':db.count_items()}


def rebuild_from_html_cache(data_dir: Path | str) -> dict:
    """Re-extract the catalog from saved HTML only. No network calls are made.

    Existing discovery rows and manual Status/Notes/title/category/thumbnail overrides are
    preserved.  For every valid archive entry, extracted links are replaced from the
    cached HTML and the item is marked completed.  Reports, duplicate groups, and the
    searchable HTML catalog are rebuilt at the end.
    """
    data_dir=Path(data_dir); db=CatalogDB(data_dir/'state'/'catalog.db')
    archive=data_dir/'html_archive'; scanned=rebuilt=blocked=failed=0
    for url,folder in legacy.scan_archive(archive):
        scanned+=1
        try:
            page,links=legacy.extract_course_from_archive(folder)
            if legacy.is_probable_block_page(Path(page.saved_html).read_text(encoding='utf-8',errors='replace')) if page.saved_html and Path(page.saved_html).exists() else False:
                blocked+=1; continue
            iid=db.upsert_discovered({'title':page.parent_title,'url':url,'url_date':page.url_date,'image_url':'','topic_slug':'video-tutorials'})
            with db.connect() as con:
                con.execute('DELETE FROM links WHERE item_id=?',(iid,))
            db.update_processed(iid,page,links)
            rebuilt+=1
        except Exception as exc:
            failed+=1; log(data_dir,f'CACHE REBUILD warning {folder}: {exc}')
    dup=run_duplicate_checker(db)
    refresh_generated_outputs(db,data_dir)
    result={'archives_scanned':scanned,'archives_rebuilt':rebuilt,'blocked_skipped':blocked,'failed':failed,'duplicate_groups':dup['groups'],'duplicates':dup['duplicates']}
    log(data_dir,f'CACHE REBUILD complete | {result}')
    return result


def import_catalog_edits(db: CatalogDB, edits_path: Path) -> int:
    data=json.loads(Path(edits_path).read_text(encoding='utf-8-sig'))
    if not isinstance(data,dict): raise ValueError('Edits file must be a JSON object')
    n=0
    for iid,e in data.items():
        if not isinstance(e,dict) or not db.get_item(iid): continue
        db.set_manual(iid,title=e.get('title'),category=e.get('category'),status=e.get('status'),notes=e.get('notes'),thumbnail=e.get('image')); n+=1
    return n


def parse_args(argv=None):
    p=argparse.ArgumentParser(description='Downloadly Catalog Manager V30')
    p.add_argument('--data-dir',default=str(DEFAULT_DATA)); p.add_argument('--topics',default='video-tutorials')
    p.add_argument('--phase',choices=['pipeline','discover','process','resume','publish','migrate','import-edits','rebuild-cache','dedupe','scheduled'],default='pipeline')
    p.add_argument('--import-edits','--edits-file',dest='edits_file',default='')
    p.add_argument('--timeout',type=int,default=45); p.add_argument('--workers',type=int,default=3); p.add_argument('--max-batches',type=int,default=0)
    p.add_argument('--block-wait-minutes',type=float,default=3.0); p.add_argument('--max-wait-minutes',type=float,default=30.0); p.add_argument('--increasing-wait',action='store_true')
    p.add_argument('--refresh-discovery',action='store_true'); p.add_argument('--download-images',action='store_true')
    p.add_argument('--translation',choices=['disabled','argos'],default='disabled'); p.add_argument('--argos-auto-install',action='store_true')
    p.add_argument('--allowed-from',default=''); p.add_argument('--allowed-to',default=''); p.add_argument('--report-every','--report-buffer-courses',dest='report_every',type=int,default=100); p.add_argument('--report-seconds',type=int,default=300)
    p.add_argument('--clear-cookies-before-run',action='store_true'); p.add_argument('--skip-migration',action='store_true'); p.add_argument('--retry-failed',action='store_true'); p.add_argument('--new-only',action='store_true'); p.add_argument('--max-block-retries',type=int,default=1)
    # Backward-compatible accepted options from V18-V21; V23 queue processing supersedes their old control loops.
    p.add_argument('--blocked-retry-seconds',type=int,default=0); p.add_argument('--maintenance-seconds',type=int,default=0); p.add_argument('--rolling-months',type=int,default=0)
    return p.parse_args(argv)


def main(argv=None) -> int:
    args=parse_args(argv); data=Path(args.data_dir); data.mkdir(parents=True,exist_ok=True)
    db=CatalogDB(data/'state'/'catalog.db'); control=Control(data,args.allowed_from,args.allowed_to); control.clear_stale()
    recovered=db.recover_interrupted_items()
    if recovered:
        log(data,f"Recovered {recovered} item(s) left in processing state by an interrupted/force-stopped worker; returned to pending.")
    legacy.COOKIE_SNAPSHOT_PATH=data/'state'/'downloadly_cookie_snapshot.json'; legacy.CLEAR_SITE_COOKIES_ON_NEW_SESSION=True
    if args.clear_cookies_before_run and args.phase not in {'rebuild-cache','dedupe','publish','migrate','import-edits'}:
        try: legacy.COOKIE_SNAPSHOT_PATH.unlink(missing_ok=True)
        except Exception: pass
        try: legacy.clear_downloadly_cookies(legacy.get_http_session())
        except Exception: pass
        legacy.reset_http_session()
    topics=[x.strip() for x in args.topics.split(',') if x.strip() in TOPICS]
    if not topics: raise SystemExit('No valid topics selected')
    try:
        # Replace any stale legacy terminal state immediately so the GUI visibly shows
        # that a new worker/resume request actually started, even while migration scans old data.
        write_progress(data,{'status':'running','phase':args.phase,'items':db.count_items(),
                             'message':'Preparing saved state and queue...' if args.phase=='resume' else 'Starting selected job...'})
        if args.phase=='rebuild-cache':
            result=rebuild_from_html_cache(data)
            write_progress(data,{'status':'completed','phase':'rebuild-cache','items':db.count_items(),'message':f"Cache rebuild complete: {result['archives_rebuilt']}/{result['archives_scanned']} archives rebuilt"})
            return 0
        existing_before_migration=int(db.state_counts(topics).get('total',0))
        should_migrate = (not args.skip_migration) and not (args.phase in {'resume','scheduled'} and existing_before_migration>0)
        if should_migrate:
            mig=migrate_existing(db,data); log(data,f"Migration/import scan: {mig}")
        elif args.phase in {'resume','scheduled'} and existing_before_migration>0:
            log(data,f"{args.phase.title()} fast path: existing SQLite queue has {existing_before_migration} item(s); skipping legacy archive/report rescan.")
        if args.phase=='migrate':
            refresh_generated_outputs(db,data); return 0
        if args.phase=='dedupe':
            result=run_duplicate_checker(db); refresh_generated_outputs(db,data)
            log(data,f"DUPLICATE CHECK complete | {result}")
            write_progress(data,{'status':'completed','phase':'dedupe','items':db.count_items(),'duplicate_groups':result['groups'],'duplicates':result['duplicates'],'message':f"Duplicate check complete: {result['groups']} groups, {result['duplicates']} duplicate posts"})
            return 0
        if args.phase=='import-edits':
            if not args.edits_file: raise SystemExit('--edits-file required')
            n=import_catalog_edits(db,Path(args.edits_file)); log(data,f'Imported {n} catalog edits'); refresh_generated_outputs(db,data); return 0
        if args.phase in {'pipeline','discover','scheduled'}:
            for slug in topics:
                discover_topic(db,data,slug,timeout=args.timeout,max_batches=args.max_batches,
                               block_wait=max(1,int(args.block_wait_minutes*60)),increasing_wait=args.increasing_wait,
                               max_wait=max(1,int(args.max_wait_minutes*60)),control=control,refresh=args.refresh_discovery,
                               report_every=args.report_every,report_seconds=args.report_seconds,new_only=(args.phase=='scheduled' or args.new_only),max_block_retries=args.max_block_retries)
        if args.phase in {'pipeline','process','resume','scheduled'}:
            if args.phase=='resume':
                counts=db.state_counts(topics)
                pending_count=len(db.pending_items(topics=topics,include_failed=args.retry_failed))
                completed_count=int(counts.get('completed',0))
                failed_count=int(counts.get('failed',0)+counts.get('error',0))
                blocked_count=int(counts.get('blocked',0))
                selected_total=int(counts.get('total',0))
                log(data,f'RESUME READY | total={selected_total} | completed={completed_count} | queued={pending_count} | blocked={blocked_count} | failed={failed_count} | retry_failed={bool(args.retry_failed)}')
                write_progress(data,{'status':'running','phase':'resume','pending':pending_count,'completed':completed_count,
                                     'failed':failed_count,'blocked':blocked_count,'items':selected_total,
                                     'message':f'Resume ready: {completed_count} completed, {pending_count} queued'})
            process_queue(db,data,topics,timeout=args.timeout,workers=args.workers,
                          block_wait=max(1,int(args.block_wait_minutes*60)),increasing_wait=args.increasing_wait,
                          max_wait=max(1,int(args.max_wait_minutes*60)),control=control,download_images=args.download_images,
                          translation=TranslationConfig(engine=args.translation,auto_install_package=args.argos_auto_install),
                          report_every=args.report_every,report_seconds=args.report_seconds,retry_failed=args.retry_failed,max_block_retries=args.max_block_retries)
        if args.phase=='scheduled':
            dup=run_duplicate_checker(db); log(data,f"Scheduled duplicate check: {dup}")
        refresh_generated_outputs(db,data)
        write_progress(data,{'status':'completed','phase':'complete','items':db.count_items(),'message':'Selected job completed; GUI may remain open.'})
        return 0
    except KeyboardInterrupt:
        refresh_generated_outputs(db,data)
        write_progress(data,{'status':'stopped','phase':'stopped','items':db.count_items(),'message':'Stopped safely. Resume later without losing saved queue/state.'})
        return 0
    except Exception as exc:
        import traceback
        log(data,'FATAL: '+traceback.format_exc())
        write_progress(data,{'status':'error','phase':'error','message':str(exc)})
        return 1

if __name__=='__main__':
    raise SystemExit(main())

[executed on device: Murali (27a4e370-9282-4fd1-98d2-8172581061df)]